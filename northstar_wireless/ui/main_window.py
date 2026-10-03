from .responsive_controls import WrappingButton as QPushButton, WrappingCheckBox as QCheckBox
import logging
from html import escape
import random

import json

import os

import time


from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook

from openpyxl.utils import get_column_letter

from openpyxl.styles import Font, Alignment

from PySide6.QtCore import Qt, QPointF, QRectF, QTimer, QSize, QStandardPaths, QSettings

from PySide6.QtGui import (
    QAction,
    QColor,
    QFont,
    QPainter,
    QPen,
    QImage,
    QPixmap,
    QUndoStack,
    QIcon,
    QKeySequence,
)

from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QFormLayout,
    QFrame,
    QGraphicsLineItem,
    QGraphicsScene,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QMainWindow,
    QMessageBox,
    QSpinBox,
    QSplitter,
    QVBoxLayout,
    QWidget,
    QScrollArea,
    QSizePolicy,
    QFileDialog,
    QDoubleSpinBox,
    QToolBar,
    QLineEdit,
    QInputDialog,
)

from dataclasses import asdict
from ..domain.models import RadioProfile, RADIO_PROFILES, MATERIALS, LinkMetrics
from ..domain.presets import DEMO_NETWORKS
from ..services.paths import resource_path
from ..services.project_store import validate_project, save_project
from ..services.propagation import PropagationEngine
from .scene_items import (
    NodeItem,
    ObstacleItem,
    LinkItem,
    MapView,
    HeatmapItem,
    HeatLegendWidget,
    SuggestMarkerItem,
    HelpDialog,
    clamp,
    dist,
)
from .commands import (
    _CmdAddNode,
    _CmdDeleteNode,
    _CmdDeleteObstacle,
    _CmdChangeObstacleMaterial,
    _CmdChangeObstacleThickness,
)
from .radio_editor import ChangeRadioCommand, edit_radio
from .commands import ChangeLabelCommand, RenameNodeCommand


class MainWindow(QMainWindow):
    def __init__(self, enable_recovery=True):
        super().__init__()
        icon_path = resource_path("assets/branding/northstar.ico")
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))

        self.setWindowTitle("NorthStar Wireless Simulator v2.2.5")
        self.resize(1400, 800)

        # ── Core state ────────────────────────────────────────────────────────
        self.scene = QGraphicsScene(self)
        self.scene.setSceneRect(QRectF(0, 0, 900, 700))
        self._syncing_ui = False
        self.scene.selectionChanged.connect(self._on_selection_changed)
        self.network_tech: Optional[str] = None  # ✅ locked after first node

        self._router_suggestion_pos: Optional[QPointF] = None
        self._router_suggestion_marker = None

        self.undo_stack = QUndoStack(self)
        self.undo_stack.cleanChanged.connect(self._on_clean_changed)
        self.undo_stack.indexChanged.connect(self._on_undo_index_changed)
        self._loading_project = False  # guards undo + live refresh while loading

        # ── Project persistence state ─────────────────────────────────────────────
        self.current_project_path: str | None = None
        self._dirty = False
        self._persistent_edits = False
        self._last_user_save_ts = 0.0

        # A stable place for autosaves/settings (per-user)
        self.settings = QSettings("NorthStar Software Works", "Wireless Simulator")

        # Autosave file path (per-user AppData)
        appdata = QStandardPaths.writableLocation(QStandardPaths.AppDataLocation)
        os.makedirs(appdata, exist_ok=True)
        self.autosave_path = os.path.join(appdata, "autosave_rf_project.json")

        # ── Routing / Mesh ───────────────────────────────────────────────────────
        self.route_enabled = False
        self.route_overlay_items: List[QGraphicsLineItem] = []
        self.route_last_sig = None
        self.route_hop_penalty_ms = 4.0  # per hop penalty (MAC/mesh overhead)

        # ── Scale / Grid settings ───────────────────────────────────────────────
        self.px_per_meter = 10.0  # default: 10 px = 1 m  (matches your old 100px=10m)
        self.grid_spacing_m = 5.0  # grid every 5 m by default
        self.show_grid = True
        self.snap_to_grid = False

        # Let items access MainWindow settings
        self.scene._mw = self

        self.view = MapView(self.scene, self)
        self.view.setAlignment(Qt.AlignLeft | Qt.AlignTop)
        self.view.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.view.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)

        self._cursor_pending_pos: Optional[QPointF] = None
        self._cursor_timer = QTimer(self)
        self._cursor_timer.setInterval(35)  # ~28 fps updates
        self._cursor_timer.timeout.connect(self._update_cursor_readout)
        self._cursor_timer.start()

        # ── Heatmap overlay ──────────────────────────────────────────────────────────
        self.heatmap = HeatmapItem(self.scene.sceneRect())
        self.scene.addItem(self.heatmap)

        self._heat_dirty = True
        self._heat_last_sig = None

        self.nodes: List[NodeItem] = []
        self.links: List[LinkItem] = []
        self.engine = PropagationEngine()

        # Per-type used numbers: e.g. {"Scale": {1,2}, "Gateway": {1,3}}
        self._used_numbers: Dict[str, set[int]] = {}

        # ── UI Layout ─────────────────────────────────────────────────────────
        root = QWidget()
        root_layout = QHBoxLayout(root)
        root_layout.setContentsMargins(8, 8, 8, 8)
        root_layout.setSpacing(8)

        # Left: map
        map_card = QFrame()
        map_card.setObjectName("Card")
        map_layout = QVBoxLayout(map_card)
        map_layout.setContentsMargins(10, 10, 10, 10)
        map_layout.setSpacing(8)

        map_header = QHBoxLayout()
        self.lbl_mode = QLabel("Mode: Select/Move")
        self.lbl_mode.setObjectName("Dim")
        map_header.addWidget(self.lbl_mode)
        map_header.addStretch(1)

        self.btn_clear = QPushButton("New / Clear")
        self.btn_clear.clicked.connect(self._clear_all)
        map_header.addWidget(self.btn_clear)

        map_layout.addLayout(map_header)
        map_layout.addWidget(self.view, 1)

        # Right: controls + metrics
        right = QVBoxLayout()
        right.setSpacing(8)

        ctrl = self._build_controls()
        metrics = self._build_metrics()

        self.controls_card = ctrl
        self.metrics_card = metrics
        self.right_scroll = None
        right.addWidget(ctrl)
        right.addWidget(metrics)
        right.addStretch(1)

        right.setContentsMargins(0, 0, 0, 0)  # ✅ remove extra padding inside scroll
        right_wrap = QWidget()
        right_wrap.setLayout(right)

        # ✅ Make right side scrollable so controls don't compress/squish
        right_scroll = QScrollArea()
        right_scroll.setWidgetResizable(True)
        right_scroll.setFrameShape(QFrame.NoFrame)
        right_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        right_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        right_scroll.setMinimumWidth(300)

        # ✅ remove internal margins that can look like a "cut off" strip
        right_scroll.setContentsMargins(0, 0, 0, 0)
        right_scroll.viewport().setContentsMargins(0, 0, 0, 0)
        right_scroll.setStyleSheet("QScrollArea { background: transparent; }")

        right_scroll.setWidget(right_wrap)
        self.right_scroll = right_scroll

        # ✅ ensure the inner widget expands to fill scroll viewport width
        right_wrap.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)

        splitter = QSplitter(Qt.Horizontal)
        splitter.addWidget(map_card)
        splitter.addWidget(right_scroll)  # ✅ use scroll area instead of direct widget
        splitter.setStretchFactor(0, 4)
        splitter.setStretchFactor(1, 3)

        root_layout.addWidget(splitter, 1)
        self.setCentralWidget(root)
        self._enhance_ui()
        self._build_guided_workspace()

        self._build_menu()
        self._build_toolbar()
        view_menu = self._view_menu
        fit = QAction("Fit Layout", self)
        fit.setShortcut("Ctrl+0")
        fit.triggered.connect(
            lambda: self.view.fitInView(
                self.scene.itemsBoundingRect().adjusted(-20, -20, 20, 20), Qt.KeepAspectRatio
            )
        )
        view_menu.addAction(fit)
        reset = QAction("Actual Scale", self)
        reset.setShortcut("Ctrl+1")
        reset.triggered.connect(self.view.resetTransform)
        view_menu.addAction(reset)

        # ── Timer: live update links & metrics ────────────────────────────────
        self.tick = QTimer(self)
        self.tick.setInterval(150)
        self.tick.timeout.connect(self._update_sim)
        self.tick.start()
        self.autosave_timer = QTimer(self)
        self.autosave_timer.setInterval(30000)
        self.autosave_timer.timeout.connect(self._autosave_tick)
        self.autosave_timer.start()

        # Seed demo
        self._add_demo()
        self.undo_stack.setClean()
        self._set_dirty(False)
        if enable_recovery:
            QTimer.singleShot(0, self._maybe_offer_recovery)
        QTimer.singleShot(50, lambda: self._on_view_resized(self.view.viewport().size()))

    # ─────────────────────────────────────────────────────────────────────────
    # UI Builders
    # ─────────────────────────────────────────────────────────────────────────

    def _help_quick_start(self):
        HelpDialog(
            self,
            title="Quick start — NorthStar Wireless Simulator",
            html="<h2>Build a layout in five stages</h2><ol><li><b>Network &amp; devices:</b> Load a demo or start a blank project. Choose a radio family and add devices.</li><li><b>Map &amp; obstacles:</b> Set the physical map scale, draw walls, and drag devices. Turn drawing off before moving items.</li><li><b>Coverage &amp; calibration:</b> Show a heatmap from a chosen transmitter. Select one device to edit its radio. Turn random fading off for repeatable readings.</li><li><b>Routes &amp; placement:</b> Select different endpoints for abstract routes and explore a relay suggestion. Real hardware must support every relay role.</li><li><b>Review, save &amp; export:</b> Review the displayed link margins, save the editable JSON project, then export PNG or Excel.</li></ol><p>Use the stage selector to jump anywhere. Back/Next only change the visible controls. Hide step hints if you prefer direct editing. Ctrl+Z undoes edits; Ctrl+S saves; Ctrl+0 fits the layout. Ctrl+click selects multiple devices.</p><p>All RF results are planning estimates. Success/retry/latency values are heuristic curves, not measurements. Read Help → Full User Guide for formulas and assumptions.</p>",
        ).show()

    def _icon(self, rel_path: str) -> QIcon:
        """
        Load an icon from assets. Includes a fallback for the common
        one-file double-folder add-data mistake (assets/icons/icons/...).
        """
        p1 = resource_path(rel_path)
        if os.path.exists(p1):
            return QIcon(p1)

        # fallback if user bundled assets/icons -> assets/icons (causes icons/icons)
        if rel_path.startswith("assets/icons/"):
            p2 = resource_path(rel_path.replace("assets/icons/", "assets/icons/icons/", 1))
            if os.path.exists(p2):
                return QIcon(p2)

        return QIcon()

    def _on_undo_index_changed(self, *_):
        self._sync_dirty_from_undo()

    def closeEvent(self, event):
        # ── Determine dirty state safely ────────────────────────────────────────
        dirty = bool(getattr(self, "_dirty", False))
        try:
            if hasattr(self, "undo_stack") and self.undo_stack is not None:
                dirty = dirty or (not self.undo_stack.isClean())
        except RuntimeError:
            # undo_stack C++ object may already be tearing down
            pass

        # ── Prompt user if needed ───────────────────────────────────────────────
        if dirty:
            res = QMessageBox.question(
                self,
                "Unsaved Changes",
                "You have unsaved changes.\n\nSave before exiting?",
                QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel,
                QMessageBox.Yes,
            )

            if res == QMessageBox.Cancel:
                event.ignore()
                return

            if res == QMessageBox.Yes:
                try:
                    if getattr(self, "current_project_path", None):
                        # Save to existing path (must exist in your codebase)
                        self._save_project_to_path(self.current_project_path)

                        # Mark clean
                        try:
                            if hasattr(self, "undo_stack") and self.undo_stack is not None:
                                self.undo_stack.setClean()
                        except RuntimeError:
                            pass

                        if hasattr(self, "_set_dirty"):
                            self._set_dirty(False)

                    else:
                        # Save As dialog
                        before = bool(getattr(self, "_dirty", False))
                        self._save_project_dialog()

                        # If still dirty, user canceled or save failed
                        if bool(getattr(self, "_dirty", False)) and before:
                            event.ignore()
                            return

                except Exception as e:
                    QMessageBox.warning(self, "Save Failed", f"Could not save project:\n\n{e}")
                    event.ignore()
                    return

        if getattr(self, "_closed_cleanly", False):
            event.accept()
            return
        self._closed_cleanly = True
        # ── Disconnect signals to avoid late callbacks during teardown ──────────
        try:
            if hasattr(self, "undo_stack") and self.undo_stack is not None:
                try:
                    self.undo_stack.indexChanged.disconnect(self._on_undo_index_changed)
                except Exception:
                    pass
                try:
                    self.undo_stack.cleanChanged.disconnect(self._on_clean_changed)
                except Exception:
                    pass
        except Exception:
            pass

        self.tick.stop()
        self._cursor_timer.stop()
        if hasattr(self, "autosave_timer"):
            self.autosave_timer.stop()
        self.settings.setValue("geometry", self.saveGeometry())
        # ── Accept close ────────────────────────────────────────────────────────
        event.accept()
        super().closeEvent(event)

    def _maybe_offer_recovery(self):
        try:
            if not os.path.exists(self.autosave_path):
                return

            # If not dirty and we just launched, still allow recovery if autosave exists
            ts = self.settings.value("last_autosave_ts", 0, type=float) or 0.0
            if ts <= 0:
                ts = os.path.getmtime(self.autosave_path)

            # If user last saved AFTER autosave, recovery usually not needed
            if self._last_user_save_ts and self._last_user_save_ts >= ts:
                return

            # If autosave is extremely old (optional cutoff)
            # if time.time() - ts > 7 * 24 * 3600: return

            when = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(ts))
            msg = (
                "A recovery autosave was found.\n\n"
                f"Autosave time: {when}\n\n"
                "Restore this autosave now?"
            )

            res = QMessageBox.question(
                self, "Crash Recovery", msg, QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes
            )

            if res != QMessageBox.Yes:
                return

            with open(self.autosave_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            data = validate_project(data)
            # Load without undo spam
            self.undo_stack.clear()
            self._apply_project_dict(data)

            # After recovery, mark dirty so user knows to Save
            self.current_project_path = None
            self._set_dirty(True)
            self.undo_stack.clear()
            self._update_window_title()

            QMessageBox.information(
                self,
                "Recovered",
                "Autosave restored.\n\nPlease Save Project to keep it permanently.",
            )

        except Exception:
            # If autosave corrupted, optionally delete it
            # os.remove(self.autosave_path)
            pass

    def _set_dirty(self, on: bool):
        if not on:
            self._persistent_edits = False
        if self._dirty == bool(on):
            return
        self._dirty = bool(on)
        self._update_window_title()

    def _sync_dirty_from_undo(self):
        # During teardown Qt can delete C++ objects before Python finishes callbacks
        try:
            us = getattr(self, "undo_stack", None)
            if us is None:
                return
            clean = us.isClean()
        except RuntimeError:
            return

        self._set_dirty(not clean or self._persistent_edits)

    def _on_clean_changed(self, clean: bool):
        # clean=True => no unsaved changes
        self._set_dirty(not clean or self._persistent_edits)

    def _update_window_title(self):
        star = "*" if self._dirty else ""
        name = (
            os.path.basename(self.current_project_path) if self.current_project_path else "Untitled"
        )
        self.setWindowTitle(f"NorthStar Wireless Simulator 2.2.5{star} — {name}")

    def _save_project_dialog(self):
        # Default name/path
        default_name = "rf_project.json"
        start_dir = ""

        try:
            last = (
                self.settings.value("last_project_path", "", type=str)
                if hasattr(self, "settings")
                else ""
            )
            if last:
                start_dir = os.path.dirname(last)
                default_name = os.path.basename(last) or default_name
        except Exception:
            pass

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Save Project",
            os.path.join(start_dir, default_name) if start_dir else default_name,
            "JSON (*.json)",
        )
        if not path:
            return
        if not path.lower().endswith(".json"):
            path += ".json"

        try:
            data = self._project_to_dict()
            data["_meta"] = {
                "version": int(data.get("version", 1)),
                "saved_ts": time.time(),
                "autosave": False,
            }

            # Atomic write (safer against crashes)
            save_project(path, data)
            if os.path.exists(self.autosave_path):
                self._remove_recovery_file()

            # Update app state
            self.current_project_path = path
            if hasattr(self, "settings"):
                self.settings.setValue("last_project_path", path)

            self._last_user_save_ts = time.time()

            # ✅ Mark as clean (no star)
            if hasattr(self, "undo_stack"):
                self.undo_stack.setClean()
            if hasattr(self, "_set_dirty"):
                self._set_dirty(False)
            if hasattr(self, "_update_window_title"):
                self._update_window_title()

            QMessageBox.information(self, "Save Project", f"Saved:\n{path}")

        except Exception as e:
            QMessageBox.warning(self, "Save Project", f"Failed:\n{e}")

    def _load_project_dialog(self):
        start_dir = ""
        try:
            last = (
                self.settings.value("last_project_path", "", type=str)
                if hasattr(self, "settings")
                else ""
            )
            if last:
                start_dir = os.path.dirname(last)
        except Exception:
            pass

        path, _ = QFileDialog.getOpenFileName(self, "Load Project", start_dir, "JSON (*.json)")
        if not path:
            return

        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)

            data = validate_project(data)
            if not self._confirm_replace():
                return
            # Load without undo spam
            if hasattr(self, "undo_stack"):
                self.undo_stack.clear()

            self._apply_project_dict(data)

            # Update app state
            self.current_project_path = path
            if hasattr(self, "settings"):
                self.settings.setValue("last_project_path", path)

            # After load: clean state
            if hasattr(self, "undo_stack"):
                self.undo_stack.clear()
                self.undo_stack.setClean()
            if hasattr(self, "_set_dirty"):
                self._set_dirty(False)
            if hasattr(self, "_update_window_title"):
                self._update_window_title()

            # Track save timestamp if file has meta
            try:
                meta = data.get("_meta") or {}
                saved_ts = float(meta.get("saved_ts", 0.0) or 0.0)
                if saved_ts > 0:
                    self._last_user_save_ts = saved_ts
            except Exception:
                pass

            QMessageBox.information(self, "Load Project", f"Loaded:\n{path}")

        except Exception as e:
            QMessageBox.warning(self, "Load Project", f"Failed:\n{e}")

    def _autosave_tick(self):
        # Don’t autosave while loading, or if there’s nothing to save
        if getattr(self, "_loading_project", False):
            return
        if len(self.nodes) == 0 and len(self._get_obstacles()) == 0:
            return

        # Optional: only autosave if dirty
        if not self._dirty:
            return

        try:
            data = self._project_to_dict()
            data["_meta"] = {
                "autosave": True,
                "timestamp": time.time(),
                "source_project": self.current_project_path or "",
            }

            tmp = self.autosave_path + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2)
            os.replace(tmp, self.autosave_path)  # atomic replace on Windows

            self.settings.setValue("last_autosave_path", self.autosave_path)
            self.settings.setValue("last_autosave_ts", float(data["_meta"]["timestamp"]))

        except Exception:
            logging.getLogger(__name__).exception("Autosave failed")
            self.statusBar().showMessage("Recovery save failed. Use File > Save Project.")

    def _project_to_dict(self) -> dict:
        # Heatmap settings (if UI exists)
        heat = {}
        if hasattr(self, "chk_heatmap"):
            heat = {
                "show": bool(self.chk_heatmap.isChecked()),
                "mode": self.cb_heat_mode.currentText()
                if hasattr(self, "cb_heat_mode")
                else "RSSI (dBm)",
                "tx": self.cb_heat_tx.currentText() if hasattr(self, "cb_heat_tx") else "",
            }

        data = {
            "version": 2,
            "simulation": {
                "extra_loss_db": self.engine.extra_loss_db,
                "live_fading": self.engine.live_fading,
            },
            "network_tech": self.network_tech,
            "scale": {
                "px_per_meter": float(getattr(self, "px_per_meter", 10.0)),
                "grid_spacing_m": float(getattr(self, "grid_spacing_m", 5.0)),
                "show_grid": bool(getattr(self, "show_grid", True)),
                "snap_to_grid": bool(getattr(self, "snap_to_grid", False)),
            },
            "links": {
                "mode": self.sp_link_mode.currentText()
                if hasattr(self, "sp_link_mode")
                else "All-to-All",
                "nn_k": int(self.sp_nn_k.value()) if hasattr(self, "sp_nn_k") else 2,
            },
            "heatmap": heat,
            "routing": {
                "enabled": self.chk_route.isChecked(),
                "mode": self.cb_route_mode.currentText(),
                "source": self.cb_route_src.currentText(),
                "destination": self.cb_route_dst.currentText(),
            },
            "nodes": [],
            "obstacles": [],
        }

        # Nodes
        for n in self.nodes:
            p = n.pos()
            data["nodes"].append(
                {
                    "name": n.name,
                    "custom_name": n.custom_name,
                    "base_label": getattr(n, "base_label", ""),
                    "type_number": int(getattr(n, "type_number", 0)),
                    "profile_key": getattr(n, "profile_key", ""),
                    "radio": asdict(n.profile),
                    "x": float(p.x()),
                    "y": float(p.y()),
                }
            )

        # Obstacles
        for o in self._get_obstacles():
            r = o.rect()  # local (0,0,w,h)
            pos = o.pos()  # scene top-left
            data["obstacles"].append(
                {
                    "material": o.material.name,
                    "thickness": float(getattr(o, "thickness", 1.0)),
                    "x": float(pos.x()),
                    "y": float(pos.y()),
                    "w": float(r.width()),
                    "h": float(r.height()),
                }
            )

        return data

    def _apply_project_dict(self, data: dict):
        """
        Clears current layout and loads everything from dict.
        Does not push undo commands.
        """
        data = validate_project(data)
        self._loading_project = True
        try:
            self.undo_stack.clear()
            self._reset_scene_references()
            # hard reset without undo
            self.scene.clear()
            self.nodes.clear()
            self.links.clear()
            self._used_numbers = {}
            self.network_tech = data.get("network_tech", None)

            self.chk_route.setChecked(False)
            # scale/grid
            sc = data.get("scale", {}) or {}
            self.px_per_meter = float(sc.get("px_per_meter", 10.0))
            self.grid_spacing_m = float(sc.get("grid_spacing_m", 5.0))
            self.show_grid = bool(sc.get("show_grid", True))
            self.snap_to_grid = bool(sc.get("snap_to_grid", False))

            self.sp_px_per_m.setValue(self.px_per_meter)
            self.sp_grid_m.setValue(self.grid_spacing_m)
            self.chk_grid.setChecked(self.show_grid)
            self.chk_snap.setChecked(self.snap_to_grid)
            # links
            lk = data.get("links", {}) or {}
            if hasattr(self, "sp_link_mode"):
                self.sp_link_mode.setCurrentText(lk.get("mode", "All-to-All"))
            if hasattr(self, "sp_nn_k"):
                self.sp_nn_k.setValue(int(lk.get("nn_k", 2)))

            # recreate heatmap item (scene.clear nukes it)
            self._ensure_heatmap_item()

            # restore nodes
            for nd in data.get("nodes", []) or []:
                prof_key = nd.get("profile_key", "")
                prof = RadioProfile(**nd["radio"]) if "radio" in nd else RADIO_PROFILES[prof_key]

                base = str(nd.get("base_label", "Node"))
                num = int(nd.get("type_number", 0)) or self._alloc_type_number(base)

                # ensure _used_numbers reflects loaded numbering
                used = self._used_numbers.setdefault(base, set())
                used.add(num)

                name = str(nd.get("name", f"{base}{num}"))
                node = NodeItem(
                    name=name,
                    profile=prof,
                    base_label=base,
                    type_number=num,
                    profile_key=prof_key if prof_key in RADIO_PROFILES else prof.name,
                    pos=QPointF(float(nd.get("x", 100)), float(nd.get("y", 100))),
                )
                node.custom_name = nd.get("custom_name", node.custom_name)
                self.scene.addItem(node)
                self.nodes.append(node)

            # restore obstacles
            for od in data.get("obstacles", []) or []:
                mat_name = str(od.get("material", "Concrete"))
                mat = MATERIALS.get(mat_name, MATERIALS["Concrete"])
                th = float(od.get("thickness", 1.0))
                rect = QRectF(
                    float(od.get("x", 200)),
                    float(od.get("y", 200)),
                    float(od.get("w", 120)),
                    float(od.get("h", 120)),
                )
                obs = ObstacleItem(rect, mat, thickness=th)
                self.scene.addItem(obs)

            sim = data.get("simulation", {})
            self.sp_extra_loss.setValue(sim.get("extra_loss_db", 0.0))
            self.chk_live_fading.setChecked(sim.get("live_fading", True))
            # UI sync
            self._refresh_profile_dropdown()
            self._update_network_status_label()
            self._refresh_heat_tx_list()
            if hasattr(self, "_refresh_route_lists"):
                self._refresh_route_lists()

            # heatmap settings
            hm = data.get("heatmap", {}) or {}
            if hasattr(self, "chk_heatmap"):
                self.chk_heatmap.setChecked(bool(hm.get("show", False)))
            if hasattr(self, "cb_heat_mode"):
                self.cb_heat_mode.setCurrentText(hm.get("mode", "RSSI (dBm)"))
            if hasattr(self, "cb_heat_tx") and hm.get("tx"):
                self.cb_heat_tx.setCurrentText(hm["tx"])

            route = data.get("routing", {})
            self.cb_route_mode.setCurrentText(route.get("mode", "Best Path (Max-Min Margin)"))
            self.cb_route_src.setCurrentText(route.get("source", ""))
            self.cb_route_dst.setCurrentText(route.get("destination", ""))
            self.chk_route.setChecked(route.get("enabled", False))
            # rebuild links + refresh visuals
            self._rebuild_links()
            self._mark_heatmap_dirty()
            self.view.viewport().update()
        finally:
            self._loading_project = False

    def _on_obstacle_thickness_changed(self, v: float):
        if self._syncing_ui:
            return

        v = float(max(0.1, v))

        # If an obstacle is selected, update it live
        sel = self.scene.selectedItems()

        for it in sel:
            if isinstance(it, ObstacleItem):
                old = float(getattr(it, "thickness", 1.0))
                new = float(v)
                if abs(old - new) > 1e-9 and not self._loading_project:
                    self.undo_stack.push(_CmdChangeObstacleThickness(self, it, old, new))
                else:
                    it.thickness = new
                    it._set_label_text()
                    self._mark_heatmap_dirty()
                break

    def _snap_point(self, p: QPointF) -> QPointF:
        if not self.snap_to_grid:
            return p
        step_px = max(1.0, self.grid_spacing_m * self.px_per_meter)
        x = round(p.x() / step_px) * step_px
        y = round(p.y() / step_px) * step_px
        return QPointF(x, y)

    def _m_per_px(self) -> float:
        return 1.0 / max(1e-6, float(self.px_per_meter))

    def _accept_router_suggestion(self):
        p = self._router_suggestion_pos
        if p is None:
            return

        prof_key = self._router_profile_key_for_network()
        prof = RADIO_PROFILES.get(prof_key)
        if prof is None:
            return

        # Enforce network tech lock
        if self.network_tech is None:
            self.network_tech = prof.tech
            self._refresh_profile_dropdown()
            self._update_network_status_label()
        elif prof.tech != self.network_tech:
            # pick a profile matching current network tech
            for k, pr in RADIO_PROFILES.items():
                if pr.tech == self.network_tech:
                    prof_key = k
                    prof = pr
                    break

        # Create a real Router node at suggested location
        base = "Router"
        num = self._alloc_type_number(base)
        name = f"{base}{num}"

        node = NodeItem(
            name=name, profile=prof, base_label=base, type_number=num, profile_key=prof_key, pos=p
        )

        self.scene.addItem(node)
        self.nodes.append(node)

        # Cleanup UI + marker
        self._reject_router_suggestion()

        # Refresh links/heatmap/node lists
        self._rebuild_links()
        self._refresh_heat_tx_list()
        self._mark_heatmap_dirty()

    def _reject_router_suggestion(self):
        self._router_suggestion_pos = None

        if self._router_suggestion_marker is not None:
            try:
                self.scene.removeItem(self._router_suggestion_marker)
            except Exception:
                pass
            self._router_suggestion_marker = None

        self.btn_accept_router.setVisible(False)
        self.btn_reject_router.setVisible(False)

    def _candidate_region_rect(self) -> QRectF:
        """
        Area we sample for router placement.
        For v1.2: use the scene rect with a margin.
        """
        sr = self.scene.sceneRect()
        margin = 40
        return sr.adjusted(margin, margin, -margin, -margin)

    def _router_profile_key_for_network(self) -> str:
        """
        Pick a reasonable profile key for the router based on active network tech.
        """
        # If unlocked, default to currently selected profile
        if self.network_tech is None:
            return self.cb_profile.currentText()

        # Prefer a profile in RADIO_PROFILES matching tech
        for k, prof in RADIO_PROFILES.items():
            if prof.tech == self.network_tech:
                return k

        # fallback
        return self.cb_profile.currentText()

    def _score_router_pos(self, p: QPointF, obstacles: list[ObstacleItem], prof_key: str) -> float:
        """
        Score candidate router position by maximizing the minimum margin
        between router and all existing nodes (both directions).
        Higher is better.
        """
        prof = RADIO_PROFILES[prof_key]

        # If no nodes, no score
        if not self.nodes:
            return -1e9

        mpp = self._m_per_px()  # ✅ meters-per-pixel from your UI scale

        # Build a temporary lightweight "tx node" proxy:
        # We reuse compute_to_point by faking a NodeItem-like object with scene_center/profile.
        class _TxProxy:
            def __init__(self, pos: QPointF, profile, name="Router?"):
                self._pos = pos
                self.profile = profile
                self.name = name

            def scene_center(self):
                return self._pos

        router_tx = _TxProxy(p, prof, "Router?")

        min_margin = 1e9

        for n in self.nodes:
            # router -> node
            m1 = self.engine.compute_to_point(router_tx, n.scene_center(), obstacles, mpp)

            # node -> router (use node's profile)
            node_tx = _TxProxy(n.scene_center(), n.profile, n.name)
            m2 = self.engine.compute_to_point(node_tx, p, obstacles, mpp)

            # worst of the two directions for this neighbor
            pair_min = min(m1.margin_db, m2.margin_db)

            # global minimum
            if pair_min < min_margin:
                min_margin = pair_min

        return float(min_margin)

    def _suggest_router_position(self):
        """
        Sample candidate positions and pick the one with best min-margin score.
        """
        if len(self.nodes) < 2:
            QMessageBox.information(self, "Suggest Router", "Add at least 2 nodes first.")
            return

        # If network is locked, ensure we're suggesting within the same tech
        prof_key = self._router_profile_key_for_network()
        prof = RADIO_PROFILES.get(prof_key)
        if prof is None:
            return

        if self.network_tech is not None and prof.tech != self.network_tech:
            # pick first matching profile
            for k, pr in RADIO_PROFILES.items():
                if pr.tech == self.network_tech:
                    prof_key = k
                    prof = pr
                    break

        obstacles = self._get_obstacles()

        region = self._candidate_region_rect()

        # Coarse sampling grid
        step = 40  # smaller = more accurate, slower
        best_score = -1e9
        best_p = None

        y = region.top()
        while y <= region.bottom():
            x = region.left()
            while x <= region.right():
                p = QPointF(x, y)

                # Small skip: avoid placing inside obstacles
                blocked = False
                for obs in obstacles:
                    rect = obs.mapRectToScene(obs.rect())  # QRectF
                    if rect.contains(p):
                        blocked = True
                        break
                if not blocked:
                    s = self._score_router_pos(p, obstacles, prof_key)
                    if s > best_score:
                        best_score = s
                        best_p = p

                x += step
            y += step

        if best_p is None:
            QMessageBox.warning(self, "Suggest Router", "No valid placement found.")
            return

        # Store suggestion + show marker + show accept/reject
        self._router_suggestion_pos = best_p
        self._show_router_suggestion_marker(best_p, best_score)

    def _show_router_suggestion_marker(self, p: QPointF, score: float):
        # Remove old marker
        if self._router_suggestion_marker is not None:
            try:
                self.scene.removeItem(self._router_suggestion_marker)
            except Exception:
                pass
            self._router_suggestion_marker = None

        self._router_suggestion_marker = SuggestMarkerItem(p)
        self.scene.addItem(self._router_suggestion_marker)

        # Buttons visible
        self.btn_accept_router.setVisible(True)
        self.btn_reject_router.setVisible(True)

        # Status line (if you have one)
        if hasattr(self, "_set_status"):
            self._set_status(f"Suggested router position: min margin ≈ {score:.1f} dB")

    def _fmt_link_tooltip(self, lk, m) -> str:
        return (
            f"<b>{escape(lk.a.name)}</b> → <b>{escape(lk.b.name)}</b> (directed estimate)<br>"
            f"RSSI: <b>{m.rssi_dbm:.1f} dBm</b><br>"
            f"Margin: <b>{m.margin_db:.1f} dB</b><br>"
            f"Success: <b>{m.success_pct:.0f}%</b><br>"
            f"Distance: <b>{m.distance_m:.1f} m</b><br>"
            f"Obstacle loss: <b>{m.obstacle_db:.0f} dB</b>"
        )

    def _on_map_mouse_moved(self, scene_pos: QPointF):
        # Only care if heatmap UI exists
        if not hasattr(self, "heat_legend"):
            return
        self._cursor_pending_pos = scene_pos

    def _update_cursor_readout(self):
        if not hasattr(self, "chk_heatmap") or not hasattr(self, "heat_legend"):
            return

        # If heatmap isn't shown, don't spam compute
        if not self.chk_heatmap.isChecked():
            self.heat_legend.clear_cursor()
            self._cursor_pending_pos = None
            return

        if self._cursor_pending_pos is None:
            return

        tx = self._get_heat_tx_node()
        if tx is None:
            self.heat_legend.set_cursor_value("No TX")
            self._cursor_pending_pos = None
            return

        obstacles = self._get_obstacles()
        p = self._cursor_pending_pos
        self._cursor_pending_pos = None

        mpp = self._m_per_px()
        m = self.engine.compute_to_point(tx, p, obstacles, mpp)
        mode = self.cb_heat_mode.currentText()

        if mode.startswith("RSSI"):
            self.heat_legend.set_cursor_value(f"{m.rssi_dbm:.1f} dBm")
        elif mode.startswith("Margin"):
            self.heat_legend.set_cursor_value(f"{m.margin_db:.1f} dB")
        else:
            self.heat_legend.set_cursor_value(f"{m.success_pct:.0f} %")

    def _on_view_resized(self, size: QSize):
        """
        Expand the scene and heatmap when the map viewport grows.
        """
        # Expand scene rect to match view size (keeping origin at 0,0)
        new_rect = QRectF(0, 0, size.width(), size.height())

        if (
            new_rect.width() > self.scene.sceneRect().width()
            or new_rect.height() > self.scene.sceneRect().height()
        ):
            self.scene.setSceneRect(new_rect)

            # Heatmap must be recreated/rescaled
            self._ensure_heatmap_item()
            self.heatmap._scene_rect = self.scene.sceneRect()
            self._mark_heatmap_dirty()

    def _on_heatmap_toggled(self, on: bool):
        self._ensure_heatmap_item()
        self.heatmap.setVisible(on)
        self._mark_heatmap_dirty()

    def _mark_heatmap_dirty(self):
        self._heat_dirty = True

    def _refresh_heat_tx_list(self):
        """
        Keep TX dropdown synced with available nodes.
        """
        cur = self.cb_heat_tx.currentText() if hasattr(self, "cb_heat_tx") else ""
        names = [n.name for n in self.nodes]

        self._syncing_ui = True
        try:
            self.cb_heat_tx.clear()
            self.cb_heat_tx.addItems(names)
            if cur in names:
                self.cb_heat_tx.setCurrentText(cur)
            elif names:
                self.cb_heat_tx.setCurrentIndex(0)
        finally:
            self._syncing_ui = False

    def _heat_use_selected_tx(self):
        sel = [it for it in self.scene.selectedItems() if isinstance(it, NodeItem)]
        if not sel:
            return
        tx = sel[0]
        self.cb_heat_tx.setCurrentText(tx.name)
        self._mark_heatmap_dirty()

    def _get_heat_tx_node(self) -> Optional[NodeItem]:
        name = self.cb_heat_tx.currentText().strip()
        for n in self.nodes:
            if n.name == name:
                return n
        return self.nodes[0] if self.nodes else None

    def _heatmap_signature(self) -> tuple:
        """
        A lightweight signature to detect layout changes.
        """
        # nodes: name + pos
        n_sig = tuple(
            sorted(
                (
                    n.name,
                    round(n.pos().x(), 1),
                    round(n.pos().y(), 1),
                    n.profile_key,
                    tuple(asdict(n.profile).values()),
                )
                for n in self.nodes
            )
        )

        # obstacles: material + pos + size (local rect)
        obs = [it for it in self.scene.items() if isinstance(it, ObstacleItem)]
        o_sig = tuple(
            sorted(
                (
                    o.material.name,
                    round(o.pos().x(), 1),
                    round(o.pos().y(), 1),
                    round(o.rect().width(), 1),
                    round(o.rect().height(), 1),
                    round(o.thickness, 2),
                )
                for o in obs
            )
        )

        tx = self.cb_heat_tx.currentText()
        mode = self.cb_heat_mode.currentText()
        show = self.chk_heatmap.isChecked()

        return (self.px_per_meter, self.engine.extra_loss_db, show, mode, tx, n_sig, o_sig)

    def _build_heatmap_image(self, tx: NodeItem, obstacles: List[ObstacleItem]) -> QImage:
        sr = self.scene.sceneRect()
        scale = 12  # pixels per sample cell; adjust for speed/quality
        w = max(10, int(sr.width() / scale))
        h = max(10, int(sr.height() / scale))

        img = QImage(w, h, QImage.Format_ARGB32)
        img.fill(0)

        mode = self.cb_heat_mode.currentText()

        for y in range(h):
            for x in range(w):
                # sample at cell center in scene coords
                px = sr.left() + (x + 0.5) * scale
                py = sr.top() + (y + 0.5) * scale
                mpp = self._m_per_px()
                m = self.engine.compute_to_point(tx, QPointF(px, py), obstacles, mpp)

                if mode.startswith("RSSI"):
                    val = m.rssi_dbm  # e.g. -110..-40
                    t = clamp((val - (-110.0)) / (-40.0 - (-110.0)), 0.0, 1.0)
                elif mode.startswith("Margin"):
                    val = m.margin_db  # e.g. -20..+20
                    t = clamp((val - (-20.0)) / (20.0 - (-20.0)), 0.0, 1.0)
                else:
                    val = m.success_pct  # 0..100
                    t = clamp(val / 100.0, 0.0, 1.0)

                # green (good) <-> red (bad)
                r = int(255 * (1.0 - t))
                g = int(255 * t)
                b = 0
                a = 90  # transparency

                img.setPixel(x, y, QColor(r, g, b, a).rgba())

        return img

    def _alloc_type_number(self, base: str) -> int:
        used = self._used_numbers.setdefault(base, set())
        n = 1
        while n in used or any(node.name == f"{base}{n}" for node in self.nodes):
            n += 1
        used.add(n)
        return n

    def _release_type_number(self, base: str, n: int):
        used = self._used_numbers.get(base)
        if used and n in used:
            used.remove(n)

    def _on_selection_changed(self):
        """
        If an obstacle is selected, reflect its material in the dropdown.
        If a node is selected, reflect its type + radio profile in the dropdowns.
        """
        if self._syncing_ui:
            return

        sel = self.scene.selectedItems()

        # Prefer obstacle sync if an obstacle is selected
        for it in sel:
            if isinstance(it, ObstacleItem):
                self._syncing_ui = True
                try:
                    self.cb_material.setCurrentText(it.material.name)
                    if hasattr(self, "sp_obs_thickness"):
                        self.sp_obs_thickness.setValue(float(getattr(it, "thickness", 1.0)))
                finally:
                    self._syncing_ui = False
                return

        # Otherwise sync node UI if a node is selected
        for it in sel:
            if isinstance(it, NodeItem):
                self._syncing_ui = True
                try:
                    # Update device label dropdown
                    self.cb_node_name.setCurrentText(it.base_label)

                    # Update profile dropdown
                    # (use the stored key so it matches exactly)
                    self.cb_profile.setCurrentText(it.profile_key)
                finally:
                    self._syncing_ui = False
                return

    def _on_heat_mode_changed(self, _=None):
        mode = self.cb_heat_mode.currentText()

        if mode.startswith("RSSI"):
            self.heat_legend.set_mode_and_range("RSSI", -110, -40, " dBm")
        elif mode.startswith("Margin"):
            self.heat_legend.set_mode_and_range("Margin", -20, 20, " dB")
        else:
            self.heat_legend.set_mode_and_range("Success", 0, 100, " %")

        self.heat_legend.clear_cursor()
        self._mark_heatmap_dirty()

    def _build_menu(self):
        bar = self.menuBar()

        # Small helper: load PNG icons that you bundle via resource_path()
        # Folder suggestion: assets/icons/*.png
        def ico(rel: str):
            p = resource_path(rel)
            return QIcon(p) if os.path.exists(p) else QIcon()

        # ─────────────────────────────────────────────────────────────────────
        # File Menu
        # ─────────────────────────────────────────────────────────────────────
        m_file = bar.addMenu("File")

        act_new = QAction(ico("assets/icons/ico_new.png"), "New Project", self)
        act_new.setShortcut(QKeySequence.New)
        act_new.triggered.connect(self._new_project_wizard)
        m_file.addAction(act_new)

        m_file.addSeparator()

        act_load = QAction(ico("assets/icons/ico_open.png"), "Load Project (JSON)...", self)
        act_load.setShortcut(QKeySequence.Open)
        act_load.triggered.connect(self._load_project_dialog)
        m_file.addAction(act_load)

        act_save = QAction(ico("assets/icons/ico_save.png"), "Save Project", self)
        act_save.setShortcut(QKeySequence.Save)
        act_save.triggered.connect(self._save_current_project)
        m_file.addAction(act_save)

        act_save_as = QAction(ico("assets/icons/ico_save.png"), "Save Project As...", self)
        act_save_as.setShortcut(QKeySequence.SaveAs)
        act_save_as.triggered.connect(self._save_project_dialog)
        m_file.addAction(act_save_as)

        # Optional: recovery restore hook if you add it later
        if hasattr(self, "_restore_autosave_prompt"):
            m_file.addSeparator()
            act_restore = QAction(ico("assets/icons/ico_open.png"), "Restore Autosave...", self)
            act_restore.triggered.connect(self._restore_autosave_prompt)
            m_file.addAction(act_restore)

        m_file.addSeparator()

        act_quit = QAction("Quit", self)
        act_quit.setShortcut(QKeySequence.Quit)
        act_quit.triggered.connect(self.close)
        m_file.addAction(act_quit)

        # ─────────────────────────────────────────────────────────────────────
        # Edit Menu (Undo / Redo)
        # ─────────────────────────────────────────────────────────────────────
        m_edit = bar.addMenu("Edit")

        if hasattr(self, "undo_stack") and self.undo_stack is not None:
            act_undo = self.undo_stack.createUndoAction(self, "Undo")
            act_undo.setShortcut(QKeySequence.Undo)
            m_edit.addAction(act_undo)

            act_redo = self.undo_stack.createRedoAction(self, "Redo")
            act_redo.setShortcut(QKeySequence.Redo)
            m_edit.addAction(act_redo)

        # ─────────────────────────────────────────────────────────────────────
        # Export Menu (optional but professional)
        # ─────────────────────────────────────────────────────────────────────
        m_export = bar.addMenu("Export")

        act_png = QAction(ico("assets/icons/ico_export.png"), "Snapshot PNG...", self)
        act_png.triggered.connect(self._export_snapshot_png)
        m_export.addAction(act_png)

        act_xlsx = QAction(ico("assets/icons/ico_export.png"), "Simulation XLSX...", self)
        act_xlsx.triggered.connect(self._export_snapshot_xlsx)
        m_export.addAction(act_xlsx)

        # ─────────────────────────────────────────────────────────────────────
        # View Menu (heatmap toggle)
        # ─────────────────────────────────────────────────────────────────────
        m_view = bar.addMenu("View")
        self._view_menu = m_view

        act_heat = QAction(ico("assets/icons/ico_heatmap.png"), "Heatmap Overlay", self)
        act_heat.setCheckable(True)
        act_heat.setToolTip("Toggle Heatmap Overlay")
        act_heat.toggled.connect(lambda on: self.chk_heatmap.setChecked(on))
        m_view.addAction(act_heat)

        # keep menu toggle synced with checkbox
        if hasattr(self, "chk_heatmap"):
            self.chk_heatmap.toggled.connect(act_heat.setChecked)

        # ─────────────────────────────────────────────────────────────────────
        # Tools Menu (routing)
        # ─────────────────────────────────────────────────────────────────────
        m_tools = bar.addMenu("Tools")

        if hasattr(self, "_compute_and_show_route"):
            act_route = QAction(ico("assets/icons/ico_route.png"), "Compute Route", self)
            act_route.setToolTip("Compute & highlight best route between selected endpoints")
            act_route.triggered.connect(self._compute_and_show_route)
            m_tools.addAction(act_route)

        # ─────────────────────────────────────────────────────────────────────
        # Help Menu
        # ─────────────────────────────────────────────────────────────────────
        m_help = bar.addMenu("Help")

        if hasattr(self, "_help_quick_start"):
            act_quick = QAction("Quick Start", self)
            act_quick.setShortcut("F1")
            act_quick.triggered.connect(self._help_quick_start)
            m_help.addAction(act_quick)

            m_help.addSeparator()

        act_guide = QAction("Full User Guide", self)
        act_guide.triggered.connect(self._open_user_guide)
        m_help.addAction(act_guide)
        act_about = QAction("About", self)
        act_about.triggered.connect(self._about)
        m_help.addAction(act_about)

    def _build_toolbar(self):
        tb = QToolBar("Main")
        tb.setMovable(False)
        tb.setFloatable(False)
        tb.setIconSize(QSize(20, 20))
        self.addToolBar(tb)

        # Helper
        def ico(name: str) -> QIcon:
            return self._icon(f"assets/icons/{name}")

        # Actions
        act_new = QAction(ico("ico_new.png"), "New", self)
        act_new.setShortcut(QKeySequence.New)
        act_new.triggered.connect(self._new_project_wizard)
        tb.addAction(act_new)

        act_open = QAction(ico("ico_open.png"), "Open", self)
        act_open.setShortcut(QKeySequence.Open)
        act_open.triggered.connect(self._load_project_dialog)
        tb.addAction(act_open)

        act_save = QAction(ico("ico_save.png"), "Save", self)
        act_save.setShortcut(QKeySequence.Save)
        act_save.triggered.connect(self._save_current_project)
        tb.addAction(act_save)

        tb.addSeparator()

        act_export_png = QAction(ico("ico_export.png"), "Export PNG", self)
        act_export_png.triggered.connect(self._export_snapshot_png)
        tb.addAction(act_export_png)

        act_export_xlsx = QAction(ico("ico_export.png"), "Export XLSX", self)
        act_export_xlsx.triggered.connect(self._export_snapshot_xlsx)
        tb.addAction(act_export_xlsx)

        tb.addSeparator()

        # Heatmap toggle (syncs with your checkbox)
        act_heat = QAction(ico("ico_heatmap.png"), "Heatmap", self)
        act_heat.setCheckable(True)
        act_heat.setToolTip("Toggle Heatmap Overlay")
        act_heat.toggled.connect(lambda on: self.chk_heatmap.setChecked(on))
        tb.addAction(act_heat)

        # Keep toolbar toggle synced if user clicks the checkbox
        self.chk_heatmap.toggled.connect(act_heat.setChecked)

        # Route compute (button-style)
        if hasattr(self, "_compute_and_show_route"):
            act_route = QAction(ico("ico_route.png"), "Compute Route", self)
            act_route.setToolTip("Compute & highlight best route between selected endpoints")
            act_route.triggered.connect(self._compute_and_show_route)
            tb.addAction(act_route)

    def _build_controls(self) -> QWidget:
        card = QFrame()
        card.setObjectName("Card")
        lay = QVBoxLayout(card)
        lay.setContentsMargins(12, 12, 12, 12)
        lay.setSpacing(10)

        title = QLabel("Controls")
        title.setFont(QFont("Segoe UI", 12, QFont.Bold))
        lay.addWidget(title)

        # ── Network / Demo group ─────────────────────────────────────────────────────
        grp_net = QGroupBox("Network / Demo")
        f0 = QFormLayout(grp_net)

        self.cb_demo_network = QComboBox()
        self.cb_demo_network.addItems(list(DEMO_NETWORKS))
        f0.addRow("Demo network:", self.cb_demo_network)

        self.btn_load_demo = QPushButton("Load Demo")
        self.btn_load_demo.clicked.connect(self._load_demo_selected)
        f0.addRow("", self.btn_load_demo)

        self.lbl_network_lock = QLabel("Active network: (none)")
        self.lbl_network_lock.setObjectName("Dim")
        f0.addRow("Status:", self.lbl_network_lock)

        self.btn_export_png = QPushButton("Export Snapshot PNG")
        self.btn_export_png.clicked.connect(self._export_snapshot_png)
        lay.addWidget(self.btn_export_png)

        self.btn_export_xlsx = QPushButton("Export XLSX (All Data)")
        self.btn_export_xlsx.clicked.connect(self._export_snapshot_xlsx)
        lay.addWidget(self.btn_export_xlsx)

        lay.addWidget(grp_net)

        # ── Add node group ───────────────────────────────────────────────────
        grp_node = QGroupBox("Add Device Node")
        f = QFormLayout(grp_node)

        self.cb_profile = QComboBox()
        self.cb_profile.addItems(list(RADIO_PROFILES.keys()))
        f.addRow("Radio profile:", self.cb_profile)
        self.cb_profile.currentTextChanged.connect(self._on_profile_changed)

        self.cb_node_name = QComboBox()
        self.cb_node_name.addItems(
            [
                "Scale",
                "Router",
                "Gateway",
                "Sensor",
                "Repeater",
                "Access Point",
                "Laptop",
                "Phone",
                "Bluetooth Peripheral",
                "BLE Beacon",
                "Headset",
                "IoT Device",
                "Wi-Fi Module",
                "Bluetooth Module",
                "BLE Module",
                "Zigbee Module",
                "LoRa Module",
                "RF Module",
                "Development Board",
                "Transceiver",
                "Test Node",
                "Custom Device",
            ]
        )
        f.addRow("Device type:", self.cb_node_name)
        self.cb_node_name.setToolTip(
            "Device type is a label. The radio profile controls the wireless technology."
        )
        self.ed_device_name = QLineEdit()
        self.ed_device_name.setMaxLength(120)
        self.ed_device_name.setPlaceholderText("Optional project name; blank = automatic")
        self.ed_device_name.setToolTip(
            "Unique name for the next device you add. Select a device and use Rename to change an existing name."
        )
        f.addRow("New device name:", self.ed_device_name)
        self.btn_rename_node = QPushButton("Rename Selected…")
        self.btn_rename_node.clicked.connect(self._rename_selected_node)
        f.addRow("", self.btn_rename_node)
        self.cb_node_name.currentTextChanged.connect(self._on_node_label_changed)

        self.btn_add_node = QPushButton("Add Node")
        self.btn_add_node.clicked.connect(self._add_node)
        f.addRow("", self.btn_add_node)

        self.btn_suggest_router = QPushButton("Suggest Router Position")
        self.btn_suggest_router.clicked.connect(self._suggest_router_position)
        lay.addWidget(self.btn_suggest_router)

        row = QHBoxLayout()
        self.btn_accept_router = QPushButton("Accept Suggestion")
        self.btn_reject_router = QPushButton("Reject")
        self.btn_accept_router.clicked.connect(self._accept_router_suggestion)
        self.btn_reject_router.clicked.connect(self._reject_router_suggestion)
        row.addWidget(self.btn_accept_router)
        row.addWidget(self.btn_reject_router)

        # hidden until we have a suggestion
        self.btn_accept_router.setVisible(False)
        self.btn_reject_router.setVisible(False)
        lay.addLayout(row)
        lay.addWidget(grp_node)

        # ── Obstacle drawing group ───────────────────────────────────────────
        grp_obs = QGroupBox("Obstacles")
        f2 = QFormLayout(grp_obs)

        self.cb_material = QComboBox()
        self.cb_material.addItems(list(MATERIALS.keys()))
        self.cb_material.currentTextChanged.connect(self._on_material_changed)
        f2.addRow("Obstacle type:", self.cb_material)
        self.cb_material.setToolTip(
            "Illustrative loss per crossing. Use thickness and measured calibration for your experiment."
        )

        self.btn_draw_obs = QPushButton("Draw Obstacles (OFF)")
        self.btn_draw_obs.setCheckable(True)
        self.btn_draw_obs.toggled.connect(self._toggle_draw_obstacles)
        f2.addRow("", self.btn_draw_obs)

        self.btn_delete_sel = QPushButton("Delete Selected")
        self.btn_delete_sel.clicked.connect(self._delete_selected)
        f2.addRow("", self.btn_delete_sel)

        lay.addWidget(grp_obs)

        grp_scale = QGroupBox("Scale / Grid")
        fs = QFormLayout(grp_scale)

        self.sp_px_per_m = QDoubleSpinBox()
        self.sp_px_per_m.setRange(1.0, 200.0)
        self.sp_px_per_m.setDecimals(2)
        self.sp_px_per_m.setValue(self.px_per_meter)
        self.sp_px_per_m.valueChanged.connect(
            lambda v: (
                setattr(self, "px_per_meter", float(v)),
                self._mark_heatmap_dirty(),
                self.view.viewport().update(),
            )
        )
        fs.addRow("Pixels per meter:", self.sp_px_per_m)

        self.sp_grid_m = QDoubleSpinBox()
        self.sp_grid_m.setRange(0.5, 100.0)
        self.sp_grid_m.setDecimals(2)
        self.sp_grid_m.setValue(self.grid_spacing_m)
        self.sp_grid_m.valueChanged.connect(
            lambda v: (setattr(self, "grid_spacing_m", float(v)), self.view.viewport().update())
        )
        fs.addRow("Grid spacing (m):", self.sp_grid_m)

        self.chk_grid = QCheckBox("Show grid")
        self.chk_grid.setChecked(self.show_grid)
        self.chk_grid.toggled.connect(
            lambda on: (setattr(self, "show_grid", bool(on)), self.view.viewport().update())
        )
        fs.addRow("", self.chk_grid)

        self.chk_snap = QCheckBox("Snap to grid")
        self.chk_snap.setChecked(self.snap_to_grid)
        self.chk_snap.toggled.connect(lambda on: setattr(self, "snap_to_grid", bool(on)))
        fs.addRow("", self.chk_snap)

        lay.addWidget(grp_scale)

        self.sp_obs_thickness = QDoubleSpinBox()
        self.sp_obs_thickness.setRange(0.1, 10.0)
        self.sp_obs_thickness.setDecimals(2)
        self.sp_obs_thickness.setSingleStep(0.25)
        self.sp_obs_thickness.setValue(1.0)
        self.sp_obs_thickness.valueChanged.connect(self._on_obstacle_thickness_changed)
        f2.addRow("Thickness (x):", self.sp_obs_thickness)

        # ── Heatmap group ────────────────────────────────────────────────────────────
        grp_heat = QGroupBox("Heatmap Overlay")
        fh = QFormLayout(grp_heat)

        self.chk_heatmap = QCheckBox("Show Heatmap")
        self.chk_heatmap.toggled.connect(self._on_heatmap_toggled)
        fh.addRow("", self.chk_heatmap)

        self.cb_heat_mode = QComboBox()
        self.cb_heat_mode.addItems(["RSSI (dBm)", "Margin (dB)", "Success (%)"])
        self.cb_heat_mode.currentTextChanged.connect(self._on_heat_mode_changed)
        fh.addRow("Mode:", self.cb_heat_mode)

        self.cb_heat_tx = QComboBox()
        self.cb_heat_tx.currentTextChanged.connect(lambda _: self._mark_heatmap_dirty())
        fh.addRow("TX Node:", self.cb_heat_tx)

        self.btn_heat_use_selected = QPushButton("Use Selected as TX")
        self.btn_heat_use_selected.clicked.connect(self._heat_use_selected_tx)
        fh.addRow("", self.btn_heat_use_selected)

        self.heat_legend = HeatLegendWidget()
        fh.addRow(self.heat_legend)

        lay.addWidget(grp_heat)

        # ── Routing / Mesh ────────────────────────────────────────────────────────
        grp_route = QGroupBox("Routing / Mesh")
        fr = QFormLayout(grp_route)

        self.chk_route = QCheckBox("Enable multi-hop routing")
        self.chk_route.toggled.connect(self._on_route_toggled)
        fr.addRow("", self.chk_route)

        self.cb_route_mode = QComboBox()
        self.cb_route_mode.addItems(["Best Path (Max-Min Margin)", "Best Path (Min Total Loss)"])
        self.cb_route_mode.currentTextChanged.connect(lambda _: self._mark_route_dirty())
        fr.addRow("Mode:", self.cb_route_mode)

        self.cb_route_src = QComboBox()
        self.cb_route_dst = QComboBox()
        self.cb_route_src.currentTextChanged.connect(lambda _: self._mark_route_dirty())
        self.cb_route_dst.currentTextChanged.connect(lambda _: self._mark_route_dirty())
        fr.addRow("Source:", self.cb_route_src)
        fr.addRow("Dest:", self.cb_route_dst)

        btns = QVBoxLayout()
        self.btn_route_use_selected = QPushButton("Use Selected (2 nodes)")
        self.btn_route_use_selected.clicked.connect(self._route_use_selected_nodes)
        btns.addWidget(self.btn_route_use_selected)

        self.btn_route_compute = QPushButton("Compute Route")
        self.btn_route_compute.clicked.connect(self._compute_and_show_route)
        btns.addWidget(self.btn_route_compute)
        fr.addRow("", btns)

        # Quick route summary label
        self.lbl_route_summary = QLabel("Route: —")
        self.lbl_route_summary.setWordWrap(True)
        self.lbl_route_summary.setObjectName("Dim")
        fr.addRow(self.lbl_route_summary)

        lay.addWidget(grp_route)

        # ── Link settings ───────────────────────────────────────────────────────────
        grp_links = QGroupBox("Links")
        f3 = QFormLayout(grp_links)

        self.sp_link_mode = QComboBox()
        self.sp_link_mode.addItems(["All-to-All", "Nearest Neighbor"])
        f3.addRow("Link mode:", self.sp_link_mode)

        self.sp_nn_k = QSpinBox()
        self.sp_nn_k.setRange(1, 20)
        self.sp_nn_k.setValue(2)

        # ✅ Make sure the control can take focus and receive clicks properly
        self.sp_nn_k.setFocusPolicy(Qt.StrongFocus)

        # ✅ (Optional but recommended) Wheel should change value, not scroll the panel
        self.sp_nn_k.setKeyboardTracking(False)  # change only on commit (cleaner)
        self.sp_nn_k.wheelEvent = lambda e, sb=self.sp_nn_k: QSpinBox.wheelEvent(sb, e)

        f3.addRow("NN count (K):", self.sp_nn_k)

        # ✅ Enable/disable K based on mode
        def _sync_nn_enabled():
            is_nn = self.sp_link_mode.currentText() == "Nearest Neighbor"
            self.sp_nn_k.setEnabled(is_nn)

        _sync_nn_enabled()
        self.sp_link_mode.currentTextChanged.connect(
            lambda _: (_sync_nn_enabled(), self._rebuild_links(), self._mark_heatmap_dirty())
        )
        self.sp_nn_k.valueChanged.connect(
            lambda _: (self._rebuild_links(), self._mark_heatmap_dirty())
        )

        lay.addWidget(grp_links)
        lay.addStretch(1)
        self._on_heat_mode_changed()
        return card

    def _build_metrics(self) -> QWidget:
        card = QFrame()
        card.setObjectName("Card")
        lay = QVBoxLayout(card)
        lay.setContentsMargins(12, 12, 12, 12)
        lay.setSpacing(10)

        title = QLabel("Metrics (A → B estimate)")
        title.setFont(QFont("Segoe UI", 12, QFont.Bold))
        lay.addWidget(title)

        self.lbl_hint = QLabel(
            "Tip: Select any node and move it. Links update live.\n"
            "To focus a link, select two nodes (Ctrl+Click)."
        )
        self.lbl_hint.setObjectName("Dim")
        self.lbl_hint.setWordWrap(True)
        lay.addWidget(self.lbl_hint)

        self.metrics_grid = QFormLayout()
        self.v_rssi = QLabel("—")
        self.v_margin = QLabel("—")
        self.v_success = QLabel("—")
        self.v_retries = QLabel("—")
        self.v_latency = QLabel("—")
        self.v_distance = QLabel("—")
        self.v_obstacles = QLabel("—")

        self.metrics_grid.addRow("RSSI:", self.v_rssi)
        self.metrics_grid.addRow("Margin:", self.v_margin)
        self.metrics_grid.addRow("Success estimate:", self.v_success)
        self.metrics_grid.addRow("Retries est:", self.v_retries)
        self.metrics_grid.addRow("Latency est:", self.v_latency)
        self.metrics_grid.addRow("Distance:", self.v_distance)
        self.metrics_grid.addRow("Obstacle loss:", self.v_obstacles)

        lay.addLayout(self.metrics_grid)
        lay.addStretch(1)
        return card

    # ─────────────────────────────────────────────────────────────────────────
    # Actions
    # ─────────────────────────────────────────────────────────────────────────

    def _compute_and_show_route(self, checked=False, *, automatic=False):
        # ── Ensure routing is enabled / UI is present ──────────────────────────
        if hasattr(self, "chk_route") and not self.chk_route.isChecked():
            self.chk_route.setChecked(True)  # triggers _on_route_toggled
        self.route_enabled = True

        # ── Basic sanity ───────────────────────────────────────────────────────
        if len(self.nodes) < 2:
            self._clear_route_overlay()
            if hasattr(self, "lbl_route_summary"):
                self.lbl_route_summary.setText("Route: —")
            if not automatic:
                QMessageBox.information(self, "Compute Route", "Add at least 2 nodes first.")
            return

        # ── Make sure route dropdowns are populated/synced ─────────────────────
        # (If you forgot to call _refresh_route_lists() after demo/add/delete, this saves you.)
        if hasattr(self, "_refresh_route_lists"):
            try:
                if hasattr(self, "cb_route_src") and self.cb_route_src.count() == 0:
                    self._refresh_route_lists()
                if hasattr(self, "cb_route_dst") and self.cb_route_dst.count() == 0:
                    self._refresh_route_lists()
            except Exception:
                pass

        if not hasattr(self, "cb_route_src") or not hasattr(self, "cb_route_dst"):
            QMessageBox.warning(self, "Compute Route", "Routing UI is not initialized.")
            return

        src_name = self.cb_route_src.currentText().strip()
        dst_name = self.cb_route_dst.currentText().strip()

        src = self._get_node_by_name(src_name)
        dst = self._get_node_by_name(dst_name)

        if src is None or dst is None:
            self._clear_route_overlay()
            if hasattr(self, "lbl_route_summary"):
                self.lbl_route_summary.setText("Route: Select different nodes")
            if not automatic:
                QMessageBox.information(
                    self,
                    "Compute Route",
                    "Please select a valid Source and Destination.\n\n"
                    "Tip: Use 'Use Selected (2 nodes)' after selecting two nodes on the map.",
                )
            return

        if src == dst:
            self._clear_route_overlay()
            if hasattr(self, "lbl_route_summary"):
                self.lbl_route_summary.setText("Route: Select different nodes")
            if not automatic:
                QMessageBox.information(
                    self, "Compute Route", "Source and Destination must be different."
                )
            return

        # ── Build routing graph ────────────────────────────────────────────────
        obstacles = self._get_obstacles()
        adj = self._build_routing_edges(obstacles)

        # ── Choose routing mode ────────────────────────────────────────────────
        mode = self.cb_route_mode.currentText() if hasattr(self, "cb_route_mode") else ""
        if "Max-Min" in mode:
            path = self._route_widest_path(src, dst, adj)
        else:
            path = self._route_min_loss(src, dst, adj)

        if not path or len(path) < 2:
            self._clear_route_overlay()
            if hasattr(self, "lbl_route_summary"):
                self.lbl_route_summary.setText(f"Route: No path ({src.name} → {dst.name})")
            if not automatic:
                QMessageBox.information(
                    self,
                    "Compute Route",
                    f"No route found from {src.name} to {dst.name}.\n\n"
                    "Try:\n"
                    "- Add a Router node\n"
                    "- Move nodes closer\n"
                    "- Reduce obstacle loss (material/thickness)\n"
                    "- Switch routing mode",
                )
            return

        # ── Compute end-to-end metrics along directed hops ──────────────────────
        mpp = self._m_per_px()
        hop_ms = float(getattr(self, "route_hop_penalty_ms", 4.0))

        e2e_latency = 0.0
        e2e_success = 1.0
        bottleneck_margin = 1e9

        hop_metrics: List[LinkMetrics] = []
        for i in range(len(path) - 1):
            a = path[i]
            b = path[i + 1]

            m = self.engine.compute_link_det(a, b, obstacles, mpp)
            hop_metrics.append(m)

            e2e_latency += float(m.latency_ms) + hop_ms
            e2e_success *= clamp(float(m.success_pct) / 100.0, 0.0, 1.0)
            bottleneck_margin = min(bottleneck_margin, float(m.margin_db))

        e2e_success_pct = e2e_success * 100.0
        hops = len(path) - 1

        # ── Draw overlay + update summary ──────────────────────────────────────
        self._draw_route_overlay(path, hop_metrics)

        path_txt = " → ".join(n.name for n in path)
        if hasattr(self, "lbl_route_summary"):
            self.lbl_route_summary.setText(
                f"Route: {path_txt}\n"
                f"Hops: {hops} | Bottleneck margin: {bottleneck_margin:.1f} dB | "
                f"E2E success: {e2e_success_pct:.0f}% | E2E latency: {e2e_latency:.1f} ms"
            )

        # Sometimes helps immediately show overlay in some Qt update states
        try:
            self.view.viewport().update()
        except Exception:
            pass

    def _draw_route_overlay(self, path: List[NodeItem], hop_metrics: List[LinkMetrics]):
        # Remove old overlay
        self._clear_route_overlay()

        # Draw thick line segments (above links, below nodes/obstacles labels)
        for i in range(len(path) - 1):
            a = path[i].scene_center()
            b = path[i + 1].scene_center()

            seg = QGraphicsLineItem(a.x(), a.y(), b.x(), b.y())
            seg.setZValue(9)

            # Color by hop success (same thresholds as your links)
            m = hop_metrics[i]
            col = (
                QColor("#2fdc71")
                if m.success_pct >= 85
                else (QColor("#f0d54e") if m.success_pct >= 55 else QColor("#ff5a6c"))
            )

            pen = QPen(col, 5)
            pen.setCapStyle(Qt.RoundCap)
            pen.setJoinStyle(Qt.RoundJoin)
            seg.setPen(pen)

            seg.setToolTip(
                f"<b>Route Hop {i + 1}</b><br>"
                f"{path[i].name} → {path[i + 1].name}<br>"
                f"Margin: <b>{m.margin_db:.1f} dB</b><br>"
                f"Success: <b>{m.success_pct:.0f}%</b><br>"
                f"Latency: <b>{m.latency_ms:.1f} ms</b><br>"
                f"Obstacle loss: <b>{m.obstacle_db:.0f} dB</b>"
            )

            self.scene.addItem(seg)
            self.route_overlay_items.append(seg)

    def _route_min_loss(self, src: NodeItem, dst: NodeItem, adj) -> Optional[List[NodeItem]]:
        """
        Dijkstra: minimize sum(loss_db) where loss_db = pathloss + obstacle loss.
        """
        import heapq
        from itertools import count

        sequence = count()

        dist_db: Dict[NodeItem, float] = {n: 1e18 for n in self.nodes}
        prev: Dict[NodeItem, Optional[NodeItem]] = {n: None for n in self.nodes}

        dist_db[src] = 0.0
        heap = [(0.0, next(sequence), src)]

        while heap:
            d, _, u = heapq.heappop(heap)
            if d != dist_db[u]:
                continue
            if u == dst:
                break

            for v, _m, loss_db in adj[u]:
                nd = d + float(loss_db)
                if nd < dist_db[v]:
                    dist_db[v] = nd
                    prev[v] = u
                    heapq.heappush(heap, (nd, next(sequence), v))

        if prev[dst] is None and dst != src:
            return None

        path = []
        cur = dst
        while cur is not None:
            path.append(cur)
            cur = prev[cur]
        path.reverse()
        return path

    def _route_widest_path(self, src: NodeItem, dst: NodeItem, adj) -> Optional[List[NodeItem]]:
        """
        Widest path: maximize the minimum margin along the path.
        """
        import heapq
        from itertools import count

        sequence = count()

        best: Dict[NodeItem, float] = {n: -1e9 for n in self.nodes}
        prev: Dict[NodeItem, Optional[NodeItem]] = {n: None for n in self.nodes}

        best[src] = 1e9
        heap = [(-best[src], next(sequence), src)]  # max-heap via negative

        visited = set()

        while heap:
            neg_w, _, u = heapq.heappop(heap)
            w = -neg_w
            if u in visited:
                continue
            visited.add(u)

            if u == dst:
                break

            for v, m, _loss_db in adj[u]:
                bottleneck = min(w, float(m.margin_db))
                if bottleneck > best[v]:
                    best[v] = bottleneck
                    prev[v] = u
                    heapq.heappush(heap, (-best[v], next(sequence), v))

        if prev[dst] is None and dst != src:
            return None

        # Reconstruct
        path = []
        cur = dst
        while cur is not None:
            path.append(cur)
            cur = prev[cur]
        path.reverse()
        return path

    def _routing_signature(self) -> tuple:
        # layout + obstacle + src/dst/mode
        n_sig = tuple(
            sorted(
                (
                    n.name,
                    round(n.pos().x(), 1),
                    round(n.pos().y(), 1),
                    n.profile_key,
                    tuple(asdict(n.profile).values()),
                )
                for n in self.nodes
            )
        )
        obs = self._get_obstacles()
        o_sig = tuple(
            sorted(
                (
                    o.material.name,
                    round(o.pos().x(), 1),
                    round(o.pos().y(), 1),
                    round(o.rect().width(), 1),
                    round(o.rect().height(), 1),
                    round(float(getattr(o, "thickness", 1.0)), 2),
                )
                for o in obs
            )
        )
        src = self.cb_route_src.currentText() if hasattr(self, "cb_route_src") else ""
        dst = self.cb_route_dst.currentText() if hasattr(self, "cb_route_dst") else ""
        mode = self.cb_route_mode.currentText() if hasattr(self, "cb_route_mode") else ""
        enabled = self.route_enabled if hasattr(self, "route_enabled") else False
        return (self.px_per_meter, self.engine.extra_loss_db, enabled, mode, src, dst, n_sig, o_sig)

    def _get_node_by_name(self, name: str) -> Optional[NodeItem]:
        for n in self.nodes:
            if n.name == name:
                return n
        return None

    def _build_routing_edges(
        self, obstacles: List[ObstacleItem]
    ) -> Dict[NodeItem, List[Tuple[NodeItem, LinkMetrics, float]]]:
        """
        Returns adjacency:
        adj[a] -> list of (b, metrics_det, loss_db)
        loss_db = pathloss + obstacle_loss (lower is better)
        """
        mpp = self._m_per_px()
        adj: Dict[NodeItem, List[Tuple[NodeItem, LinkMetrics, float]]] = {n: [] for n in self.nodes}

        # Consider all pairs for routing (independent of displayed link mode)
        for i in range(len(self.nodes)):
            for j in range(len(self.nodes)):
                if i == j:
                    continue
                a = self.nodes[i]
                b = self.nodes[j]

                m = self.engine.compute_link_det(a, b, obstacles, mpp)

                # total loss in dB = tx - rssi (deterministic)
                loss_db = float(a.profile.tx_dbm - m.rssi_dbm)

                if m.margin_db >= 0:
                    adj[a].append((b, m, loss_db))

        return adj

    def _refresh_route_lists(self):
        if not hasattr(self, "cb_route_src") or not hasattr(self, "cb_route_dst"):
            return

        src_cur = self.cb_route_src.currentText()
        dst_cur = self.cb_route_dst.currentText()

        names = [n.name for n in self.nodes]

        self._syncing_ui = True
        try:
            self.cb_route_src.clear()
            self.cb_route_dst.clear()
            self.cb_route_src.addItems(names)
            self.cb_route_dst.addItems(names)

            if src_cur in names:
                self.cb_route_src.setCurrentText(src_cur)
            elif names:
                self.cb_route_src.setCurrentIndex(0)

            if dst_cur in names:
                self.cb_route_dst.setCurrentText(dst_cur)
            elif len(names) > 1:
                self.cb_route_dst.setCurrentIndex(1)
            elif names:
                self.cb_route_dst.setCurrentIndex(0)
        finally:
            self._syncing_ui = False

    def _on_route_toggled(self, on: bool):
        self.route_enabled = bool(on)
        if not on:
            self._clear_route_overlay()
            if hasattr(self, "lbl_route_summary"):
                self.lbl_route_summary.setText("Route: —")
        self._mark_route_dirty()

    def _mark_route_dirty(self):
        self.route_last_sig = None  # force recompute next tick

    def _route_use_selected_nodes(self):
        sel = [it for it in self.scene.selectedItems() if isinstance(it, NodeItem)]
        if len(sel) < 2:
            return
        a, b = sel[0], sel[1]
        self.cb_route_src.setCurrentText(a.name)
        self.cb_route_dst.setCurrentText(b.name)
        self._mark_route_dirty()
        self._compute_and_show_route()

    def _clear_route_overlay(self):
        for it in self.route_overlay_items:
            try:
                self.scene.removeItem(it)
            except Exception:
                pass
        self.route_overlay_items.clear()

    def _export_snapshot_xlsx(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Simulation Data (XLSX)", "rf_simulation.xlsx", "Excel Workbook (*.xlsx)"
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        try:
            # Ensure the sim is up to date so link tooltips/metrics reflect current layout
            self._update_sim()

            snapshot = {
                (lk.a, lk.b): self.engine.compute_link_det(
                    lk.a, lk.b, self._get_obstacles(), self._m_per_px()
                )
                for lk in self.links
            }
            wb = Workbook()

            # Helpers
            def write_header(ws, row, headers):
                bold = Font(bold=True)
                for c, h in enumerate(headers, start=1):
                    cell = ws.cell(row=row, column=c, value=h)
                    cell.font = bold
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.freeze_panes = "A2"

            def autosize(ws):
                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        v = "" if cell.value is None else str(cell.value)
                        max_len = max(max_len, len(v))
                    ws.column_dimensions[col_letter].width = min(60, max(10, max_len + 2))

            # ── Summary ──────────────────────────────────────────────────────────
            ws = wb.active
            ws.title = "Summary"
            ws.append(["Application", "NorthStar Wireless Simulator 2.2.5"])
            ws.append(
                [
                    "Model",
                    "Deterministic planning snapshot; success/retry/latency are heuristic estimates",
                ]
            )
            ws.append(["Pixels per meter", self.px_per_meter])
            ws.append(["Additional loss (dB)", self.engine.extra_loss_db])
            ws.append(["Live fading enabled", self.engine.live_fading])
            ws.append(["Network tech", self.network_tech or "(unlocked)"])
            ws.append(
                [
                    "Link mode",
                    self.sp_link_mode.currentText() if hasattr(self, "sp_link_mode") else "",
                ]
            )
            ws.append(["NN K", int(self.sp_nn_k.value()) if hasattr(self, "sp_nn_k") else ""])
            ws.append(["Heatmap visible", self.chk_heatmap.isChecked()])
            ws.append(["Heatmap metric", self.cb_heat_mode.currentText()])
            ws.append(["Heatmap transmitter", self.cb_heat_tx.currentText()])
            ws.append(["Routing enabled", self.chk_route.isChecked()])
            ws.append(["Routing mode", self.cb_route_mode.currentText()])
            ws.append(["Route source", self.cb_route_src.currentText()])
            ws.append(["Route destination", self.cb_route_dst.currentText()])
            ws.append(["Route summary", self.lbl_route_summary.text()])
            ws.append(["Nodes", len(self.nodes)])
            ws.append(["Obstacles", len(self._get_obstacles())])
            ws.append(["Links", len(self.links)])
            autosize(ws)

            # ── Nodes ────────────────────────────────────────────────────────────
            ws = wb.create_sheet("Nodes")
            headers = [
                "Name",
                "Type",
                "Type #",
                "Profile",
                "Tech",
                "TX (dBm)",
                "Sensitivity (dBm)",
                "Freq (MHz)",
                "X",
                "Y",
                "Reference loss (dB)",
                "Path loss exponent",
                "Fading sigma (dB)",
            ]
            write_header(ws, 1, headers)

            for n in self.nodes:
                p = n.pos()
                prof = n.profile
                ws.append(
                    [
                        n.name,
                        getattr(n, "base_label", ""),
                        getattr(n, "type_number", ""),
                        getattr(n, "profile_key", ""),
                        getattr(prof, "tech", ""),
                        getattr(prof, "tx_dbm", ""),
                        getattr(prof, "sens_dbm", ""),
                        getattr(prof, "freq_mhz", ""),
                        float(p.x()),
                        float(p.y()),
                        prof.pl_d0_db,
                        prof.n,
                        prof.noise_sigma_db,
                    ]
                )
            autosize(ws)

            # ── Obstacles ────────────────────────────────────────────────────────
            ws = wb.create_sheet("Obstacles")
            headers = [
                "Material",
                "Atten (dB)",
                "X",
                "Y",
                "Width",
                "Height",
                "Thickness multiplier",
            ]
            write_header(ws, 1, headers)

            for o in self._get_obstacles():
                # ObstacleItem uses local rect (0,0,w,h) + scene pos
                r = o.rect()
                pos = o.pos()
                ws.append(
                    [
                        o.material.name,
                        float(o.material.attenuation_db),
                        float(pos.x()),
                        float(pos.y()),
                        float(r.width()),
                        float(r.height()),
                        o.thickness,
                    ]
                )
            autosize(ws)

            # ── Links ────────────────────────────────────────────────────────────
            ws = wb.create_sheet("Links")
            headers = [
                "A",
                "B",
                "Distance (m)",
                "RSSI (dBm)",
                "Margin (dB)",
                "Success (%)",
                "Obstacle loss (dB)",
                "Retries est",
                "Latency (ms)",
            ]
            write_header(ws, 1, headers)

            for lk in self.links:
                m = snapshot[(lk.a, lk.b)]
                ws.append(
                    [
                        lk.a.name,
                        lk.b.name,
                        float(m.distance_m),
                        float(m.rssi_dbm),
                        float(m.margin_db),
                        float(m.success_pct),
                        float(m.obstacle_db),
                        float(m.retries_est),
                        float(m.latency_ms),
                    ]
                )
            autosize(ws)

            # ── Warnings ─────────────────────────────────────────────────────────
            ws = wb.create_sheet("Warnings")
            headers = [
                "Severity",
                "A",
                "B",
                "Margin (dB)",
                "Success (%)",
                "RSSI (dBm)",
                "Obstacle loss (dB)",
                "Distance (m)",
                "Reason",
            ]
            write_header(ws, 1, headers)

            # Thresholds (tweak as you like)
            BAD_MARGIN = 0.0  # below this = likely unreliable
            WARN_MARGIN = 6.0  # below this = caution
            BAD_SUCCESS = 70.0  # below this = poor
            WARN_SUCCESS = 85.0  # below this = caution

            def severity_for(margin_db: float, success_pct: float) -> str:
                if margin_db < BAD_MARGIN or success_pct < BAD_SUCCESS:
                    return "CRITICAL"
                if margin_db < WARN_MARGIN or success_pct < WARN_SUCCESS:
                    return "WARN"
                return ""

            for lk in self.links:
                m = snapshot[(lk.a, lk.b)]

                sev = severity_for(m.margin_db, m.success_pct)
                if not sev:
                    continue

                reasons = []
                if m.margin_db < BAD_MARGIN:
                    reasons.append(f"Negative margin ({m.margin_db:.1f} dB)")
                elif m.margin_db < WARN_MARGIN:
                    reasons.append(f"Low margin ({m.margin_db:.1f} dB)")

                if m.success_pct < BAD_SUCCESS:
                    reasons.append(f"Low success ({m.success_pct:.0f}%)")
                elif m.success_pct < WARN_SUCCESS:
                    reasons.append(f"Marginal success ({m.success_pct:.0f}%)")

                ws.append(
                    [
                        sev,
                        lk.a.name,
                        lk.b.name,
                        float(m.margin_db),
                        float(m.success_pct),
                        float(m.rssi_dbm),
                        float(m.obstacle_db),
                        float(m.distance_m),
                        "; ".join(reasons),
                    ]
                )

            autosize(ws)

            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith(
                            ("=", "+", "-", "@")
                        ):
                            cell.data_type = "s"
            wb.save(path)
            QMessageBox.information(self, "Export XLSX", f"Saved:\n{path}")

        except Exception as e:
            QMessageBox.warning(self, "Export XLSX", f"Export failed:\n{e}")

    def _export_snapshot_png(self):
        # Choose file
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Snapshot (PNG)", "rf_snapshot.png", "PNG Image (*.png)"
        )
        if not path:
            return
        if not path.lower().endswith(".png"):
            path += ".png"

        # Ensure heatmap exists if enabled (scene.clear can delete it)
        if hasattr(self, "chk_heatmap") and self.chk_heatmap.isChecked():
            self._ensure_heatmap_item()
            self._mark_heatmap_dirty()
            self._update_sim()  # force a fresh heatmap frame

        self._update_sim()
        # What to render: full scene rect
        sr = self.scene.sceneRect()

        # High-res scale factor (2.0 = 2x, 3.0 = very crisp)
        scale = 3.0

        img_w = int(sr.width() * scale)
        img_h = int(sr.height() * scale)

        image = QImage(img_w, img_h, QImage.Format_ARGB32)
        image.fill(QColor("#0f1116"))  # match your dark background

        painter = QPainter(image)
        painter.setRenderHint(QPainter.Antialiasing, True)

        # Scale up for high-res render
        painter.scale(scale, scale)

        # Render scene
        self.scene.render(painter, target=QRectF(0, 0, sr.width(), sr.height()), source=sr)
        painter.end()

        ok = image.save(path, "PNG")
        if ok:
            QMessageBox.information(self, "Export Snapshot", f"Saved:\n{path}")
        else:
            QMessageBox.warning(self, "Export Snapshot", "Failed to save PNG.")

    def _ensure_heatmap_item(self):
        """
        scene.clear() deletes QGraphicsItems (including HeatmapItem).
        Recreate it if it's missing/deleted.
        """
        try:
            # Accessing a deleted Qt object raises RuntimeError
            _ = self.heatmap.isVisible()
            return
        except Exception:
            pass

        self.heatmap = HeatmapItem(self.scene.sceneRect())
        self.scene.addItem(self.heatmap)
        self.heatmap.setVisible(
            self.chk_heatmap.isChecked() if hasattr(self, "chk_heatmap") else False
        )
        self._heat_dirty = True
        self._heat_last_sig = None

    def _load_demo_selected(self):
        if self._confirm_replace():
            self._add_demo(DEMO_NETWORKS[self.cb_demo_network.currentText()])

    def _on_node_label_changed(self, base: str):
        if self._syncing_ui:
            return
        for it in self.scene.selectedItems():
            if isinstance(it, NodeItem) and it.base_label != base:
                self.undo_stack.push(ChangeLabelCommand(self, it, base))
                break

    def _on_profile_changed(self, prof_name: str):
        if self._syncing_ui:
            return

        prof = RADIO_PROFILES.get(prof_name)
        if prof is None:
            return

        # If a node is selected, attempt to apply to it
        sel = self.scene.selectedItems()
        for it in sel:
            if not isinstance(it, NodeItem):
                continue

            # ── Enforce network lock ────────────────────────────────────────────
            if self.network_tech is not None and prof.tech != self.network_tech:
                QMessageBox.warning(
                    self,
                    "Incompatible Profile",
                    f"This layout is locked to {self.network_tech}.\n\n"
                    f"You selected a {prof.tech} profile ({prof.name}).\n\n"
                    "Clear the layout to start a different network.",
                )
                # snap UI back to the node's current profile
                self._syncing_ui = True
                try:
                    self.cb_profile.setCurrentText(it.profile_key)
                finally:
                    self._syncing_ui = False
                return

            self.undo_stack.push(ChangeRadioCommand(self, it, prof, prof_name))

        # No selected node → this only affects future nodes (still OK)
        # But if network tech is already locked, keep dropdown filtered elsewhere.

    def _about(self):
        QMessageBox.information(
            self,
            "About",
            "NorthStar Wireless Simulator 2.2.5\nNorthStar Software Works\n\nZigbee, FHSS, Wi-Fi and Bluetooth coverage planning.\n\nRSSI and margin use a simplified log-distance model.\nSuccess, retries and latency are illustrative estimates.\nRoutes are abstract graph paths, not protocol emulation.\nNo hardware is scanned or controlled.\nSee docs/User_Guide.md for assumptions and calibration.",
        )

    def _on_material_changed(self, name: str):
        # Ignore changes triggered by UI sync (selection updates)
        if self._syncing_ui:
            return

        # Update draw mode material for new obstacles
        self.view.draw_material = MATERIALS.get(name, MATERIALS["Concrete"])

        # If an obstacle is selected, update it live
        sel = self.scene.selectedItems()

        for it in sel:
            if isinstance(it, ObstacleItem):
                old = it.material
                new = self.view.draw_material
                if old != new and not self._loading_project:
                    self.undo_stack.push(_CmdChangeObstacleMaterial(self, it, old, new))
                else:
                    it.material = new
                    it._set_label_text()
                    it._apply_style()
                    self._mark_heatmap_dirty()
                break

    def _toggle_draw_obstacles(self, enabled: bool):
        self.view.set_draw_mode(enabled)
        if enabled:
            self.btn_draw_obs.setText("Draw Obstacles (ON)")
            self.lbl_mode.setText("Mode: Draw Obstacles")
        else:
            self.btn_draw_obs.setText("Draw Obstacles (OFF)")
            self.lbl_mode.setText("Mode: Select/Move")

    def _valid_device_name(self, name, exclude=None):
        if not name or len(name) > 120 or any(ord(c) < 32 or ord(c) == 127 for c in name):
            QMessageBox.warning(
                self,
                "Invalid name",
                "Use 1–120 characters without line breaks or control characters.",
            )
            return False
        if any(n is not exclude and n.name == name for n in self.nodes):
            QMessageBox.warning(self, "Name already used", "Choose a unique name for each device.")
            return False
        return True

    def _rename_selected_node(self):
        nodes = [n for n in self.scene.selectedItems() if isinstance(n, NodeItem)]
        if len(nodes) != 1:
            QMessageBox.information(
                self, "Select a device", "Select exactly one device on the map to rename it."
            )
            return
        node = nodes[0]
        name, accepted = QInputDialog.getText(
            self, "Rename device", "Project device name:", text=node.name
        )
        name = name.strip()
        if accepted and self._valid_device_name(name, exclude=node) and name != node.name:
            self.undo_stack.push(RenameNodeCommand(self, node, name))

    def _add_node(self):
        if len(self.nodes) >= 100:
            QMessageBox.information(
                self, "Device limit", "Use up to 100 devices per exploratory project."
            )
            return
        prof_key = self.cb_profile.currentText()
        prof = RADIO_PROFILES[prof_key]
        base = self.cb_node_name.currentText()

        custom_name = self.ed_device_name.text().strip()
        if custom_name and not self._valid_device_name(custom_name):
            return

        # ✅ Enforce single-network tech in a layout
        if self.network_tech is None:
            self.network_tech = prof.tech
            self._refresh_profile_dropdown()
            self._update_network_status_label()
        elif prof.tech != self.network_tech:
            QMessageBox.warning(
                self,
                "Incompatible Device",
                f"This layout is locked to {self.network_tech}.\n\n"
                f"You tried to add a {prof.tech} device ({prof.name}), which is not compatible.\n\n"
                "Clear the layout to start a different network.",
            )
            return

        num = self._alloc_type_number(base)
        name = custom_name or f"{base}{num}"

        x = random.uniform(120, 900)
        y = random.uniform(120, 560)

        node = NodeItem(
            name=name,
            profile=prof,
            base_label=base,
            type_number=num,
            profile_key=prof_key,
            pos=QPointF(x, y),
        )

        node.custom_name = bool(custom_name)
        self.ed_device_name.clear()
        if not self._loading_project:
            self.undo_stack.push(_CmdAddNode(self, node))
        else:
            self.scene.addItem(node)
            self.nodes.append(node)

        self._rebuild_links()
        self._refresh_heat_tx_list()
        if hasattr(self, "_refresh_route_lists"):
            self._refresh_route_lists()
        self._mark_heatmap_dirty()

    def _refresh_profile_dropdown(self):
        self._syncing_ui = True
        try:
            cur = self.cb_profile.currentText()
            self.cb_profile.clear()

            keys = list(RADIO_PROFILES.keys())
            if self.network_tech is not None:
                keys = [k for k in keys if RADIO_PROFILES[k].tech == self.network_tech]

            self.cb_profile.addItems(keys)

            if cur in keys:
                self.cb_profile.setCurrentText(cur)
            elif keys:
                self.cb_profile.setCurrentIndex(0)
        finally:
            self._syncing_ui = False

    def _update_network_status_label(self):
        if self.network_tech is None:
            self.lbl_network_lock.setText("Active network: (none)")
        else:
            self.lbl_network_lock.setText(f"Active network: {self.network_tech}")

    def _delete_selected(self):
        selected = self.scene.selectedItems()
        if not selected:
            return

        # ── During project load: direct delete (no undo spam) ───────────────────
        if getattr(self, "_loading_project", False):
            did_anything = False

            for it in selected:
                if isinstance(it, NodeItem):
                    self._release_type_number(it.base_label, it.type_number)
                    if it in self.nodes:
                        self.nodes.remove(it)
                    self.scene.removeItem(it)
                    did_anything = True

                elif isinstance(it, ObstacleItem):
                    self.scene.removeItem(it)
                    did_anything = True

                else:
                    # ignore transient items (links, heatmap, overlays)
                    continue

            if not did_anything:
                return

            self._rebuild_links()

            # Smart unlock if no nodes remain
            if len(self.nodes) == 0:
                self.network_tech = None
                self._refresh_profile_dropdown()
                self._update_network_status_label()

            self._refresh_heat_tx_list()
            self._mark_heatmap_dirty()

            if hasattr(self, "_refresh_route_lists"):
                self._refresh_route_lists()
            if hasattr(self, "_mark_route_dirty"):
                self._mark_route_dirty()

            return

        # ── Normal operation: use undo commands ────────────────────────────────
        did_anything = False

        # Delete nodes first (links depend on them)
        for it in selected:
            if isinstance(it, NodeItem):
                self.undo_stack.push(_CmdDeleteNode(self, it))
                did_anything = True

        # Delete obstacles next
        for it in selected:
            if isinstance(it, ObstacleItem):
                self.undo_stack.push(_CmdDeleteObstacle(self, it))
                did_anything = True

        if not did_anything:
            QMessageBox.information(self, "Delete Selected", "Select a node or obstacle to delete.")
            return

        # Safety refresh (commands should already do this, but keep in sync)
        self._rebuild_links()

        # Smart unlock if no nodes remain
        if len(self.nodes) == 0:
            self.network_tech = None
            self._refresh_profile_dropdown()
            self._update_network_status_label()

            # ⭐ This change might not be represented in undo stack clean-state,
            # so mark dirty explicitly.
            if hasattr(self, "_set_dirty"):
                self._set_dirty(True)

        self._refresh_heat_tx_list()
        self._mark_heatmap_dirty()
        if hasattr(self, "_refresh_route_lists"):
            self._refresh_route_lists()
        if hasattr(self, "_mark_route_dirty"):
            self._mark_route_dirty()

    def _clear_all(self):
        if not self._confirm_replace():
            return
        self._reset_scene_references()
        self.current_project_path = None
        self.undo_stack.clear()
        # Clear scene + reset state
        self.scene.clear()
        self.nodes.clear()
        self.links.clear()
        self._used_numbers = {}

        # Optional: clear route overlay state (if you have it)
        if hasattr(self, "_clear_route_overlay"):
            self._clear_route_overlay()

        # Unlock network tech
        self.network_tech = None

        # scene.clear() deletes items (including heatmap) — recreate
        self._ensure_heatmap_item()

        # Reset project path? (optional; keep if you want "New project")
        # self.current_project_path = None

        # Undo stack: clearing is a user action; make it a clean baseline
        if hasattr(self, "undo_stack"):
            self.undo_stack.clear()
            self.undo_stack.setClean()

        # Refresh UI
        self._refresh_profile_dropdown()
        self._update_network_status_label()
        self._refresh_heat_tx_list()
        self._mark_heatmap_dirty()

        if hasattr(self, "_refresh_route_lists"):
            self._refresh_route_lists()
        if hasattr(self, "_mark_route_dirty"):
            self._mark_route_dirty()

        # ⭐ Mark dirty (this is a user change) unless we're loading
        if not getattr(self, "_loading_project", False):
            if hasattr(self, "_set_dirty"):
                self._set_dirty(True)

    # ─────────────────────────────────────────────────────────────────────────────
    # Demo Seed
    # ─────────────────────────────────────────────────────────────────────────────

    def _add_demo(self, network_tech: str = "ZIGBEE"):
        self._reset_scene_references()
        self.current_project_path = None
        self.undo_stack.clear()
        # ── Clear scene + reset state ───────────────────────────────────────────
        try:
            self.scene.clear()
        except Exception:
            pass

        self.nodes.clear()
        self.links.clear()
        self._used_numbers = {}

        # Optional: clear route overlay state
        if hasattr(self, "_clear_route_overlay"):
            try:
                self._clear_route_overlay()
            except Exception:
                pass

        # Lock the network tech BEFORE creating nodes
        self.network_tech = network_tech
        for label, tech in DEMO_NETWORKS.items():
            if tech == network_tech:
                self.cb_demo_network.setCurrentText(label)
                break

        # scene.clear() deletes items (including heatmap) — recreate safely
        if hasattr(self, "_ensure_heatmap_item"):
            try:
                self._ensure_heatmap_item()
            except Exception:
                pass

        # Undo stack: demo should start clean (only if undo_stack exists)
        if hasattr(self, "undo_stack") and self.undo_stack is not None:
            try:
                self.undo_stack.clear()
                self.undo_stack.setClean()
            except Exception:
                pass

        # ── Helper to add fixed nodes ───────────────────────────────────────────
        def add_fixed(base: str, prof_key: str, p: QPointF):
            num = self._alloc_type_number(base)
            n = NodeItem(
                name=f"{base}{num}",
                profile=RADIO_PROFILES[prof_key],
                base_label=base,
                type_number=num,
                profile_key=prof_key,
                pos=p,
            )
            self.scene.addItem(n)
            self.nodes.append(n)

        # ── Seed demo nodes (single-tech only) ──────────────────────────────────
        key = next(k for k, p in RADIO_PROFILES.items() if p.tech == network_tech)
        add_fixed("Gateway", key, QPointF(200, 300))
        add_fixed("Sensor", key, QPointF(520, 270))
        add_fixed("Router", key, QPointF(820, 380))

        # ── Demo obstacle ───────────────────────────────────────────────────────
        try:
            obs = ObstacleItem(QRectF(390, 200, 140, 220), MATERIALS["Drywall"])
            self.scene.addItem(obs)
        except Exception:
            pass

        # ── Refresh links + UI filtering ────────────────────────────────────────
        self._rebuild_links()

        if hasattr(self, "_refresh_profile_dropdown"):
            self._refresh_profile_dropdown()

        if hasattr(self, "_update_network_status_label"):
            self._update_network_status_label()

        if hasattr(self, "_refresh_heat_tx_list"):
            self._refresh_heat_tx_list()

        if hasattr(self, "_mark_heatmap_dirty"):
            self._mark_heatmap_dirty()

        if hasattr(self, "_refresh_route_lists"):
            self._refresh_route_lists()

        if hasattr(self, "_mark_route_dirty"):
            self._mark_route_dirty()

        # ⭐ Demo modifies the current session project (only if dirty system exists)
        if not getattr(self, "_loading_project", False):
            if hasattr(self, "_set_dirty"):
                self._set_dirty(True)

    # ─────────────────────────────────────────────────────────────────────────
    # Links + Simulation
    # ─────────────────────────────────────────────────────────────────────────

    def _get_obstacles(self) -> List[ObstacleItem]:
        out: List[ObstacleItem] = []
        for it in self.scene.items():
            if isinstance(it, ObstacleItem):
                out.append(it)
        return out

    def _rebuild_links(self):
        # Remove old link items from scene
        for lk in self.links:
            self.scene.removeItem(lk)
        self.links.clear()

        if len(self.nodes) < 2:
            return

        mode = self.sp_link_mode.currentText()

        if mode == "All-to-All":
            pairs = []
            for i in range(len(self.nodes)):
                for j in range(i + 1, len(self.nodes)):
                    pairs.append((self.nodes[i], self.nodes[j]))
        else:
            # Nearest Neighbor (K)
            k = int(self.sp_nn_k.value())
            pairs = []
            for i, ni in enumerate(self.nodes):
                dlist = []
                for j, nj in enumerate(self.nodes):
                    if i == j:
                        continue
                    dlist.append((dist(ni.scene_center(), nj.scene_center()), nj))
                dlist.sort(key=lambda t: t[0])
                for _, nj in dlist[:k]:
                    # avoid duplicates
                    if (nj, ni) not in pairs:
                        pairs.append((ni, nj))

        for a, b in pairs:
            lk = LinkItem(a, b)
            self.scene.addItem(lk)
            self.links.append(lk)

    def _update_sim(self):
        self._refresh_guidance()
        # Gather obstacles once
        obstacles = self._get_obstacles()

        # ── Links: rebuild only when needed ─────────────────────────────────────
        # (Mode/K changes should call _rebuild_links() directly, but keep a cheap safety here.)
        try:
            want_nn = self.sp_link_mode.currentText() == "Nearest Neighbor"
            want_k = int(self.sp_nn_k.value())
        except Exception:
            want_nn = False
            want_k = 2

        # Cache the last settings on the instance
        if not hasattr(self, "_last_link_mode"):
            self._last_link_mode = None
            self._last_nn_k = None

        mode_key = "NN" if want_nn else "ALL"
        positions = tuple((n.pos().x(), n.pos().y()) for n in self.nodes)
        moved_nn = want_nn and positions != getattr(self, "_last_nn_positions", None)
        if self._last_link_mode != mode_key or self._last_nn_k != want_k or moved_nn:
            self._rebuild_links()
            self._last_link_mode = mode_key
            self._last_nn_k = want_k
            self._last_nn_positions = positions

        # ── Update each link geometry + color by computed metrics ───────────────
        for lk in self.links:
            a = lk.a.scene_center()
            b = lk.b.scene_center()
            lk.setLine(a.x(), a.y(), b.x(), b.y())

            mpp = self._m_per_px()
            m = self.engine.compute_link(lk.a, lk.b, obstacles, mpp)
            lk.setPen(QPen(self._link_color(m), 2))
            lk.setToolTip(self._fmt_link_tooltip(lk, m))  # ✅ hover tooltip

        # ── Update metrics panel for "focused link" ─────────────────────────────
        self._update_metrics_panel(obstacles)

        # ── Heatmap refresh (only when changed) ─────────────────────────────────
        if hasattr(self, "chk_heatmap") and self.chk_heatmap.isChecked():
            # ✅ scene.clear() may have deleted the heatmap item
            self._ensure_heatmap_item()

            sig = self._heatmap_signature()
            if self._heat_dirty or sig != self._heat_last_sig:
                tx = self._get_heat_tx_node()
                if tx is not None:
                    self.heatmap._scene_rect = self.scene.sceneRect()
                    img = self._build_heatmap_image(tx, obstacles)
                    self.heatmap.update_heatmap(img)
                    self._heat_last_sig = sig
                    self._heat_dirty = False

        # ── Routing refresh (only when changed) ───────────────────────────────────
        if hasattr(self, "chk_route") and self.route_enabled:
            sig = self._routing_signature()
            if sig != self.route_last_sig:
                self._compute_and_show_route(automatic=True)
                self.route_last_sig = sig

    def _link_color(self, m: LinkMetrics) -> QColor:
        # success-based color
        if m.success_pct >= 85:
            return QColor("#2fdc71")  # green
        if m.success_pct >= 55:
            return QColor("#f0d54e")  # yellow
        return QColor("#ff5a6c")  # red

    def _update_metrics_panel(self, obstacles: List[ObstacleItem]):
        sel = [it for it in self.scene.selectedItems() if isinstance(it, NodeItem)]
        if len(sel) >= 2:
            a, b = sel[0], sel[1]
            mpp = self._m_per_px()
            m = self.engine.compute_link(a, b, obstacles, mpp)

            self.v_rssi.setText(f"{m.rssi_dbm:.1f} dBm")
            self.v_margin.setText(f"{m.margin_db:.1f} dB")
            self.v_success.setText(f"{m.success_pct:.0f} %")
            self.v_retries.setText(f"{m.retries_est:.1f}")
            self.v_latency.setText(f"{m.latency_ms:.1f} ms")
            self.v_distance.setText(f"{m.distance_m:.1f} m")
            self.v_obstacles.setText(f"{m.obstacle_db:.1f} dB")
        else:
            self.v_rssi.setText("—")
            self.v_margin.setText("—")
            self.v_success.setText("—")
            self.v_retries.setText("—")
            self.v_latency.setText("—")
            self.v_distance.setText("—")
            self.v_obstacles.setText("—")

    def _reset_scene_references(self):
        self.route_overlay_items = []
        self.route_last_sig = None
        self._router_suggestion_marker = None
        self._router_suggestion_pos = None
        if hasattr(self, "btn_accept_router"):
            self.btn_accept_router.hide()
            self.btn_reject_router.hide()

    def _confirm_replace(self):
        if self._loading_project or not self._dirty:
            return True
        answer = QMessageBox.question(
            self,
            "Unsaved project",
            "Save your current project before replacing it?",
            QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel,
            QMessageBox.Save,
        )
        if answer == QMessageBox.Cancel:
            return False
        if answer == QMessageBox.Save:
            self._save_current_project()
            return not self._dirty
        return True

    def _save_project_to_path(self, path):
        save_project(path, self._project_to_dict())
        self._last_user_save_ts = time.time()
        self.undo_stack.setClean()
        self._set_dirty(False)
        if os.path.exists(self.autosave_path):
            self._remove_recovery_file()

    def _simulation_changed(self):
        self.engine.extra_loss_db = self.sp_extra_loss.value()
        self.engine.live_fading = self.chk_live_fading.isChecked()
        self._mark_heatmap_dirty()
        self._mark_route_dirty()
        self._setting_changed()

    def _edit_selected_radio(self):
        selected = [n for n in self.nodes if n.isSelected()]
        if len(selected) != 1:
            QMessageBox.information(self, "Edit radio", "Select one device node first.")
            return
        edit_radio(self, selected[0])

    def _model_help(self):
        HelpDialog(
            self,
            title="Model assumptions",
            html="<h2>Planning estimates</h2><p>Each project models one compatible radio family. Start a New Project to choose another family. Wi-Fi bands are separate families; BLE PHY presets share a planning family assuming compatible PHY support.</p><p>RSSI = TX − reference path loss − 10n log10(distance) − obstacle loss − additional loss + optional random fading. Reference distance is 1 meter; minimum modeled distance is 0.5 m.</p><p>Success percentage, retries and latency use generic heuristic curves and are not measurements or protocol-specific guarantees. Bluetooth pairing, Wi-Fi association, MAC contention, antenna patterns, channel overlap and spectral interference are not simulated.</p><p>Routes are abstract directed graph paths and exclude hops below receiver sensitivity. Ordinary phones, BLE beacons and headsets do not thereby become mesh relays. Set TX power and sensitivity from your device specifications and calibrate walls against measurements.</p>",
        ).show()

    def _enhance_ui(self):
        brand = QLabel()
        brand.setPixmap(
            QPixmap(resource_path("assets/branding/northstar.png")).scaled(
                46, 46, Qt.KeepAspectRatio, Qt.SmoothTransformation
            )
        )
        self.statusBar().addPermanentWidget(brand)
        self.statusBar().showMessage(
            "NorthStar Software Works • Planning estimates • One radio family per project"
        )
        grp = QGroupBox("Radio Calibration")
        self.calibration_group = grp
        form = QFormLayout(grp)
        self.sp_extra_loss = QDoubleSpinBox()
        self.sp_extra_loss.setRange(0, 100)
        self.sp_extra_loss.setSuffix(" dB")
        self.sp_extra_loss.setToolTip(
            "Additional aggregate loss for a conservative planning allowance; not a spectral interference model."
        )
        self.chk_live_fading = QCheckBox("Animate random fading")
        self.chk_live_fading.setChecked(True)
        form.addRow("Additional loss:", self.sp_extra_loss)
        form.addRow(self.chk_live_fading)
        edit = QPushButton("Edit Selected Radio…")
        edit.clicked.connect(self._edit_selected_radio)
        self.btn_edit_radio = edit
        form.addRow(edit)
        help_button = QPushButton("Model Assumptions")
        help_button.clicked.connect(self._model_help)
        form.addRow(help_button)
        # Prepend the calibration card to the scrollable controls.
        scroll = self.findChild(QScrollArea)
        scroll.widget().layout().insertWidget(0, grp)
        self.sp_extra_loss.valueChanged.connect(self._simulation_changed)
        self.chk_live_fading.toggled.connect(self._simulation_changed)
        for combo in self.findChildren(QComboBox):
            combo.setSizeAdjustPolicy(QComboBox.AdjustToMinimumContentsLengthWithIcon)
            combo.setMinimumContentsLength(8)
            combo.currentTextChanged.connect(combo.setToolTip)
            combo.setToolTip(combo.currentText())
        for label in self.findChildren(QLabel):
            label.setWordWrap(True)
        for form_layout in self.findChildren(QFormLayout):
            form_layout.setRowWrapPolicy(QFormLayout.WrapAllRows)
        for control in [self.sp_px_per_m, self.sp_grid_m, self.sp_nn_k]:
            control.valueChanged.connect(self._setting_changed)
        for control in [self.chk_grid, self.chk_snap, self.chk_heatmap, self.chk_route]:
            control.toggled.connect(self._setting_changed)
        for control in [
            self.cb_heat_mode,
            self.cb_heat_tx,
            self.cb_route_mode,
            self.cb_route_src,
            self.cb_route_dst,
            self.sp_link_mode,
        ]:
            control.currentTextChanged.connect(self._setting_changed)
        geometry = self.settings.value("geometry")
        if geometry:
            self.restoreGeometry(geometry)
        else:
            screen = QApplication.primaryScreen().availableGeometry()
            self.resize(min(1400, int(screen.width() * 0.9)), min(850, int(screen.height() * 0.9)))

    def _setting_changed(self, *_):
        if not self._loading_project and not self._syncing_ui:
            self._persistent_edits = True
            self._set_dirty(True)
        self._mark_route_dirty()

    def _build_guided_workspace(self):
        from .guided_workspace import GuidedWorkspace

        self.workflow = GuidedWorkspace(self)
        container = self.right_scroll.widget()
        layout = container.layout()
        while layout.count():
            item = layout.takeAt(0)
            if item.widget():
                item.widget().hide()
        layout.addWidget(self.workflow)
        layout.addStretch(1)
        self.controls_card.hide()
        self._refresh_guidance()

    def _refresh_guidance(self):
        if hasattr(self, "workflow"):
            self.workflow.refresh()

    def _new_project_wizard(self):
        from .guided_workspace import new_project_dialog

        new_project_dialog(self)

    def _save_current_project(self):
        if not self.current_project_path:
            self._save_project_dialog()
            self._refresh_guidance()
            return
        try:
            self._save_project_to_path(self.current_project_path)
            self.statusBar().showMessage("Project saved", 5000)
        except Exception as error:
            QMessageBox.warning(self, "Save failed", f"Your project could not be saved: {error}")
        self._refresh_guidance()

    def _remove_recovery_file(self):
        try:
            if os.path.exists(self.autosave_path):
                os.remove(self.autosave_path)
        except OSError:
            logging.getLogger(__name__).exception("Could not remove old recovery file")

    def _open_user_guide(self):
        from PySide6.QtCore import QUrl
        from PySide6.QtGui import QDesktopServices

        if not QDesktopServices.openUrl(QUrl.fromLocalFile(resource_path("docs/User_Guide.html"))):
            QMessageBox.information(
                self,
                "User guide",
                "Open docs/User_Guide.html from your project or installation folder.",
            )
